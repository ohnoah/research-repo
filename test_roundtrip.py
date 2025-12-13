#!/usr/bin/env python3
"""
Test script for docxlang codec round-trip and template usage.
"""

import json
from pathlib import Path
from docx import Document
from docx.shared import Pt
from docxlang_codec import (
    docx_to_docxlang,
    docxlang_to_docx,
    build_style_catalog_from_docx,
    iter_block_items,
)
from docx.table import Table
from docx.text.paragraph import Paragraph


TEMPLATE_PATH = "test-proposal-template.docx"
ROUNDTRIP_OUTPUT = "roundtrip_output.docx"
NEW_CONTENT_OUTPUT = "new_content_output.docx"


def print_section(title: str):
    print("\n" + "=" * 60)
    print(f" {title}")
    print("=" * 60)


def analyze_document(doc_path: str, label: str):
    """Analyze a document and return key metrics."""
    doc = Document(doc_path)

    print(f"\n--- {label} ---")

    # Count paragraphs and tables in body
    paragraphs = []
    tables = []
    for item in iter_block_items(doc):
        if isinstance(item, Paragraph):
            paragraphs.append(item)
        elif isinstance(item, Table):
            tables.append(item)

    print(f"Body paragraphs: {len(paragraphs)}")
    print(f"Body tables: {len(tables)}")
    print(f"Sections: {len(doc.sections)}")

    # Analyze styles used
    para_styles = set()
    char_styles = set()
    for p in paragraphs:
        if p.style:
            para_styles.add(p.style.name)
        for r in p.runs:
            try:
                if r.style and r.style.name and r.style.name != "Default Paragraph Font":
                    char_styles.add(r.style.name)
            except:
                pass

    table_styles = set()
    for t in tables:
        try:
            if t.style:
                table_styles.add(t.style.name)
        except:
            pass

    print(f"Paragraph styles used: {sorted(para_styles)}")
    print(f"Character styles used: {sorted(char_styles)}")
    print(f"Table styles used: {sorted(table_styles)}")

    # Check headers/footers
    for i, section in enumerate(doc.sections):
        has_header = bool(section.header.paragraphs and any(p.text.strip() for p in section.header.paragraphs))
        has_footer = bool(section.footer.paragraphs and any(p.text.strip() for p in section.footer.paragraphs))
        print(f"Section {i}: header={has_header}, footer={has_footer}")

    return {
        "paragraphs": paragraphs,
        "tables": tables,
        "para_styles": para_styles,
        "char_styles": char_styles,
        "table_styles": table_styles,
    }


def compare_paragraphs(orig_paras, rt_paras, label: str):
    """Compare paragraphs between original and round-tripped."""
    print(f"\n--- {label} Paragraph Comparison ---")

    # Compare count
    if len(orig_paras) != len(rt_paras):
        print(f"DIFFERENCE: Paragraph count differs: {len(orig_paras)} vs {len(rt_paras)}")
    else:
        print(f"Paragraph count matches: {len(orig_paras)}")

    # Compare text content
    min_len = min(len(orig_paras), len(rt_paras))
    text_diffs = []
    style_diffs = []

    for i in range(min_len):
        orig_text = orig_paras[i].text or ""
        rt_text = rt_paras[i].text or ""

        if orig_text != rt_text:
            text_diffs.append((i, orig_text[:50], rt_text[:50]))

        orig_style = orig_paras[i].style.name if orig_paras[i].style else None
        rt_style = rt_paras[i].style.name if rt_paras[i].style else None

        if orig_style != rt_style:
            style_diffs.append((i, orig_style, rt_style))

    if text_diffs:
        print(f"Text differences found ({len(text_diffs)} paragraphs):")
        for idx, orig, rt in text_diffs[:5]:  # Show first 5
            print(f"  Para {idx}: '{orig}...' vs '{rt}...'")
        if len(text_diffs) > 5:
            print(f"  ... and {len(text_diffs) - 5} more")
    else:
        print("All paragraph texts match!")

    if style_diffs:
        print(f"Style differences found ({len(style_diffs)} paragraphs):")
        for idx, orig, rt in style_diffs[:5]:
            print(f"  Para {idx}: '{orig}' vs '{rt}'")
        if len(style_diffs) > 5:
            print(f"  ... and {len(style_diffs) - 5} more")
    else:
        print("All paragraph styles match!")


def compare_run_formatting(orig_paras, rt_paras):
    """Compare run-level formatting."""
    print("\n--- Run Formatting Comparison ---")

    diffs = []
    min_len = min(len(orig_paras), len(rt_paras))

    for p_idx in range(min_len):
        orig_runs = orig_paras[p_idx].runs
        rt_runs = rt_paras[p_idx].runs

        if len(orig_runs) != len(rt_runs):
            diffs.append(f"Para {p_idx}: run count {len(orig_runs)} vs {len(rt_runs)}")
            continue

        for r_idx in range(len(orig_runs)):
            orig_r = orig_runs[r_idx]
            rt_r = rt_runs[r_idx]

            # Check bold
            if orig_r.bold != rt_r.bold:
                diffs.append(f"Para {p_idx} Run {r_idx}: bold {orig_r.bold} vs {rt_r.bold}")
            # Check italic
            if orig_r.italic != rt_r.italic:
                diffs.append(f"Para {p_idx} Run {r_idx}: italic {orig_r.italic} vs {rt_r.italic}")
            # Check font size
            orig_size = orig_r.font.size.pt if orig_r.font.size else None
            rt_size = rt_r.font.size.pt if rt_r.font.size else None
            if orig_size != rt_size:
                diffs.append(f"Para {p_idx} Run {r_idx}: font size {orig_size} vs {rt_size}")

    if diffs:
        print(f"Run formatting differences ({len(diffs)} total):")
        for d in diffs[:10]:
            print(f"  {d}")
        if len(diffs) > 10:
            print(f"  ... and {len(diffs) - 10} more")
    else:
        print("No run formatting differences detected!")


def test_roundtrip():
    """Test 1: Round-trip the template."""
    print_section("TEST 1: ROUND-TRIP")

    # Convert to docxlang
    print("\n1. Converting template to DocxLang...")
    docxlang = docx_to_docxlang(TEMPLATE_PATH, include_style_defs=True)

    # Save intermediate JSON for inspection
    with open("template_docxlang.json", "w", encoding="utf-8") as f:
        json.dump(docxlang, f, indent=2, ensure_ascii=False)
    print(f"   Saved DocxLang to template_docxlang.json")
    print(f"   Blocks: {len(docxlang['blocks'])}")
    print(f"   Paragraph styles: {docxlang['styles']['paragraph']}")
    print(f"   Character styles: {docxlang['styles']['character']}")
    print(f"   Table styles: {docxlang['styles']['table']}")

    # Convert back to DOCX
    print("\n2. Converting DocxLang back to DOCX...")
    docxlang_to_docx(
        docxlang,
        TEMPLATE_PATH,
        ROUNDTRIP_OUTPUT,
        anchor_text="<<CONTENT>>",  # Look for placeholder
        clear_body_if_no_anchor=True  # If no placeholder, replace body
    )
    print(f"   Saved round-trip output to {ROUNDTRIP_OUTPUT}")

    # Analyze both
    print("\n3. Analyzing documents...")
    orig_info = analyze_document(TEMPLATE_PATH, "Original Template")
    rt_info = analyze_document(ROUNDTRIP_OUTPUT, "Round-tripped")

    # Compare
    compare_paragraphs(orig_info["paragraphs"], rt_info["paragraphs"], "Body")
    compare_run_formatting(orig_info["paragraphs"], rt_info["paragraphs"])

    return docxlang


def test_new_content(base_docxlang: dict):
    """Test 2: Create new document with different content using template styling."""
    print_section("TEST 2: NEW CONTENT WITH TEMPLATE STYLING")

    # Create new content using the same styles from the template
    styles = base_docxlang["styles"]

    # Find some key styles to use
    para_styles = styles.get("paragraph", {})

    # Print available styles
    print("\nAvailable paragraph styles from template:")
    for code, name in para_styles.items():
        print(f"  {code}: {name}")

    # Create sample content - we'll try to use the styles from the template
    sample_blocks = []

    # Find heading and body styles
    heading_style = None
    body_style = None
    for code, name in para_styles.items():
        name_lower = name.lower()
        if "heading" in name_lower or "title" in name_lower:
            if heading_style is None or "1" in name:
                heading_style = code
        elif "normal" in name_lower or "body" in name_lower:
            body_style = code

    # Use first style if we didn't find specific ones
    if not heading_style and para_styles:
        heading_style = list(para_styles.keys())[0]
    if not body_style and para_styles:
        body_style = list(para_styles.keys())[-1] if len(para_styles) > 1 else list(para_styles.keys())[0]

    print(f"\nUsing heading style: {heading_style} ({para_styles.get(heading_style, 'N/A')})")
    print(f"Using body style: {body_style} ({para_styles.get(body_style, 'N/A')})")

    # Build sample content
    sample_blocks = [
        {
            "type": "paragraph",
            "pStyle": heading_style,
            "runs": [{"type": "text", "text": "Sample Project Proposal"}]
        },
        {
            "type": "paragraph",
            "pStyle": body_style,
            "runs": [{"type": "text", "text": "This is a sample document created using the docxlang codec. It uses the same styles as the template but with completely different content."}]
        },
        {
            "type": "paragraph",
            "pStyle": body_style,
            "runs": [
                {"type": "text", "text": "This paragraph demonstrates "},
                {"type": "text", "text": "bold text", "override": {"bold": True}},
                {"type": "text", "text": " and "},
                {"type": "text", "text": "italic text", "override": {"italic": True}},
                {"type": "text", "text": " mixed with regular text."}
            ]
        },
        {
            "type": "paragraph",
            "pStyle": heading_style,
            "runs": [{"type": "text", "text": "Project Objectives"}]
        },
        {
            "type": "paragraph",
            "pStyle": body_style,
            "runs": [{"type": "text", "text": "1. First objective - demonstrate round-trip conversion"}]
        },
        {
            "type": "paragraph",
            "pStyle": body_style,
            "runs": [{"type": "text", "text": "2. Second objective - preserve template styling"}]
        },
        {
            "type": "paragraph",
            "pStyle": body_style,
            "runs": [{"type": "text", "text": "3. Third objective - support LLM document generation"}]
        },
        {
            "type": "paragraph",
            "pStyle": heading_style,
            "runs": [{"type": "text", "text": "Conclusion"}]
        },
        {
            "type": "paragraph",
            "pStyle": body_style,
            "runs": [{"type": "text", "text": "The docxlang codec successfully demonstrates the ability to generate Word documents with consistent styling from a template."}]
        },
    ]

    # Create the new docxlang
    new_docxlang = {
        "schema": "docxlang/v1",
        "styles": styles,
        "blocks": sample_blocks,
    }

    # Save for inspection
    with open("new_content_docxlang.json", "w", encoding="utf-8") as f:
        json.dump(new_docxlang, f, indent=2, ensure_ascii=False)
    print(f"\nSaved new content DocxLang to new_content_docxlang.json")

    # Generate the document
    print("\nGenerating new document from template...")
    docxlang_to_docx(
        new_docxlang,
        TEMPLATE_PATH,
        NEW_CONTENT_OUTPUT,
        anchor_text="<<CONTENT>>",
        clear_body_if_no_anchor=True
    )
    print(f"Saved to {NEW_CONTENT_OUTPUT}")

    # Analyze
    print("\n--- Analyzing new content document ---")
    new_info = analyze_document(NEW_CONTENT_OUTPUT, "New Content Document")

    # Compare styling
    print("\n--- Style Inheritance Check ---")
    orig_info = analyze_document(TEMPLATE_PATH, "Original Template (for comparison)")

    # Check if new document uses same style names
    new_styles_used = new_info["para_styles"]
    orig_styles_available = orig_info["para_styles"]

    styles_matched = new_styles_used & orig_styles_available
    styles_missing = new_styles_used - orig_styles_available

    print(f"\nStyles used in new doc that exist in template: {sorted(styles_matched)}")
    if styles_missing:
        print(f"Styles used but not in template: {sorted(styles_missing)}")
    else:
        print("All styles in new document exist in template!")

    return new_docxlang


def detailed_style_comparison():
    """Compare style properties between documents."""
    print_section("DETAILED STYLE ANALYSIS")

    orig_doc = Document(TEMPLATE_PATH)
    rt_doc = Document(ROUNDTRIP_OUTPUT)
    new_doc = Document(NEW_CONTENT_OUTPUT)

    print("\n--- Sample Paragraph Style Properties ---")

    for doc, label in [(orig_doc, "Original"), (rt_doc, "Roundtrip"), (new_doc, "NewContent")]:
        print(f"\n{label}:")
        for p in list(iter_block_items(doc))[:3]:  # First 3 paragraphs
            if isinstance(p, Paragraph) and p.text.strip():
                style_name = p.style.name if p.style else "None"
                align = p.alignment
                space_after = p.paragraph_format.space_after
                space_after_pt = space_after.pt if space_after else None
                print(f"  '{p.text[:40]}...' style={style_name}, align={align}, space_after={space_after_pt}pt")


def main():
    print_section("DOCXLANG CODEC TEST SUITE")
    print(f"Template: {TEMPLATE_PATH}")

    # Test 1: Round-trip
    base_docxlang = test_roundtrip()

    # Test 2: New content
    test_new_content(base_docxlang)

    # Detailed comparison
    detailed_style_comparison()

    print_section("SUMMARY")
    print("""
Files created:
  - template_docxlang.json    : DocxLang representation of the template
  - roundtrip_output.docx     : Template round-tripped through codec
  - new_content_docxlang.json : Sample content in DocxLang format
  - new_content_output.docx   : New document with sample content

Please open the DOCX files in Word/LibreOffice to visually verify:
1. roundtrip_output.docx should look identical to the template
2. new_content_output.docx should have the same styling/formatting
   as the template but with completely different text content
""")


if __name__ == "__main__":
    main()
