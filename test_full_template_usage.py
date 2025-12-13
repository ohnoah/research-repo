#!/usr/bin/env python3
"""
Test using ALL template styles to create a new document.
"""

import json
from docx import Document
from docxlang_codec import docxlang_to_docx, iter_block_items
from docx.text.paragraph import Paragraph


def create_sample_proposal():
    """Create a sample proposal using all available template styles."""

    # Load the template's styles from the extracted docxlang
    with open("template_docxlang.json") as f:
        template_docxlang = json.load(f)

    styles = template_docxlang["styles"]
    style_defs = template_docxlang.get("style_defs", {}).get("paragraph", {})

    print("Available styles and their definitions:")
    for code, name in styles["paragraph"].items():
        defs = style_defs.get(code, {})
        print(f"  {code}: {name}")
        if defs:
            print(f"      Font: {defs.get('font', {})}")
            print(f"      Para: {defs.get('para', {})}")

    # Style mapping based on the template:
    # S1: Project Title (CVHP) - Blue, bold, 14pt, centered - for main title
    # S2: Document Title (CVHP) - Blue, 12pt, centered - for subtitles
    # S3: Header (CVHP) - Blue, bold - for section headers
    # S4: Main Bullet (CVHP) - for main bullet points
    # S5: Sub-bullet 1 (CVHP) - Cambria font - for first-level sub-bullets
    # S6: Sub-bullet 2 (CVHP) - Cambria font - for second-level sub-bullets

    sample_blocks = [
        # Title section
        {
            "type": "paragraph",
            "pStyle": "S1",
            "runs": [{"type": "text", "text": "AI Integration Feasibility Study"}]
        },
        {
            "type": "paragraph",
            "pStyle": "S2",
            "runs": [{"type": "text", "text": "Technical Proposal"}]
        },
        {
            "type": "paragraph",
            "pStyle": "S2",
            "runs": [{"type": "text", "text": "December 2025"}]
        },

        # Context section
        {
            "type": "paragraph",
            "pStyle": "S3",
            "runs": [{"type": "text", "text": "Executive Summary"}]
        },
        {
            "type": "paragraph",
            "pStyle": "S4",
            "runs": [{"type": "text", "text": "This proposal outlines a comprehensive approach to integrating AI capabilities into existing business processes"}]
        },
        {
            "type": "paragraph",
            "pStyle": "S4",
            "runs": [{"type": "text", "text": "The project aims to deliver measurable improvements in operational efficiency and decision-making accuracy"}]
        },
        {
            "type": "paragraph",
            "pStyle": "S4",
            "runs": [{"type": "text", "text": "Key deliverables include:"}]
        },
        {
            "type": "paragraph",
            "pStyle": "S5",
            "runs": [{"type": "text", "text": "Technical architecture documentation"}]
        },
        {
            "type": "paragraph",
            "pStyle": "S5",
            "runs": [{"type": "text", "text": "Proof-of-concept implementation"}]
        },
        {
            "type": "paragraph",
            "pStyle": "S5",
            "runs": [{"type": "text", "text": "ROI analysis and recommendations"}]
        },

        # Page break
        {"type": "page_break"},

        # Approach section
        {
            "type": "paragraph",
            "pStyle": "S3",
            "runs": [{"type": "text", "text": "Proposed Approach"}]
        },
        {
            "type": "paragraph",
            "pStyle": "S4",
            "runs": [{"type": "text", "text": "Phase 1: Discovery and Assessment"}]
        },
        {
            "type": "paragraph",
            "pStyle": "S5",
            "runs": [{"type": "text", "text": "Conduct stakeholder interviews to understand current workflows"}]
        },
        {
            "type": "paragraph",
            "pStyle": "S5",
            "runs": [{"type": "text", "text": "Analyze existing data infrastructure and integration points"}]
        },
        {
            "type": "paragraph",
            "pStyle": "S6",
            "runs": [{"type": "text", "text": "Document API endpoints and data formats"}]
        },
        {
            "type": "paragraph",
            "pStyle": "S6",
            "runs": [{"type": "text", "text": "Evaluate security and compliance requirements"}]
        },

        {
            "type": "paragraph",
            "pStyle": "S4",
            "runs": [{"type": "text", "text": "Phase 2: Design and Prototyping"}]
        },
        {
            "type": "paragraph",
            "pStyle": "S5",
            "runs": [{"type": "text", "text": "Develop technical architecture based on discovery findings"}]
        },
        {
            "type": "paragraph",
            "pStyle": "S5",
            "runs": [{"type": "text", "text": "Create working prototype demonstrating key capabilities"}]
        },

        {
            "type": "paragraph",
            "pStyle": "S4",
            "runs": [{"type": "text", "text": "Phase 3: Implementation and Rollout"}]
        },
        {
            "type": "paragraph",
            "pStyle": "S5",
            "runs": [{"type": "text", "text": "Deploy solution with phased rollout plan"}]
        },
        {
            "type": "paragraph",
            "pStyle": "S5",
            "runs": [{"type": "text", "text": "Provide training and documentation for end users"}]
        },

        # Budget section
        {
            "type": "paragraph",
            "pStyle": "S3",
            "runs": [{"type": "text", "text": "Team and Investment"}]
        },
        {
            "type": "paragraph",
            "pStyle": "S4",
            "runs": [
                {"type": "text", "text": "Project will be led by "},
                {"type": "text", "text": "Dr. Sarah Mitchell, Technical Director", "override": {"color": "EE0000"}},
                {"type": "text", "text": " with support from the AI Engineering team"}
            ]
        },
        {
            "type": "paragraph",
            "pStyle": "S4",
            "runs": [
                {"type": "text", "text": "Total Investment: $"},
                {"type": "text", "text": "450", "override": {"color": "EE0000"}},
                {"type": "text", "text": "K"}
            ]
        },
        {
            "type": "paragraph",
            "pStyle": "S5",
            "runs": [
                {"type": "text", "text": "Professional services: $"},
                {"type": "text", "text": "380", "override": {"color": "EE0000"}},
                {"type": "text", "text": "K"}
            ]
        },
        {
            "type": "paragraph",
            "pStyle": "S5",
            "runs": [
                {"type": "text", "text": "Infrastructure costs: $"},
                {"type": "text", "text": "50", "override": {"color": "EE0000"}},
                {"type": "text", "text": "K"}
            ]
        },
        {
            "type": "paragraph",
            "pStyle": "S5",
            "runs": [
                {"type": "text", "text": "Contingency: $"},
                {"type": "text", "text": "20", "override": {"color": "EE0000"}},
                {"type": "text", "text": "K"}
            ]
        },
    ]

    new_docxlang = {
        "schema": "docxlang/v1",
        "styles": styles,
        "blocks": sample_blocks,
    }

    # Save the docxlang
    with open("full_sample_docxlang.json", "w", encoding="utf-8") as f:
        json.dump(new_docxlang, f, indent=2)
    print("\nSaved full_sample_docxlang.json")

    # Generate the document
    output_path = "full_sample_output.docx"
    docxlang_to_docx(
        new_docxlang,
        "test-proposal-template.docx",
        output_path,
        anchor_text="<<CONTENT>>",
        clear_body_if_no_anchor=True
    )
    print(f"Generated {output_path}")

    return output_path


def analyze_output(doc_path: str):
    """Analyze the generated document."""
    print(f"\n{'=' * 60}")
    print(f" ANALYSIS OF {doc_path}")
    print("=" * 60)

    doc = Document(doc_path)

    # Check sections
    print(f"\nSections: {len(doc.sections)}")
    for i, section in enumerate(doc.sections):
        header_text = " | ".join([p.text for p in section.header.paragraphs if p.text.strip()])
        footer_text = " | ".join([p.text for p in section.footer.paragraphs if p.text.strip()])
        print(f"  Section {i}: header='{header_text[:50]}...', footer='{footer_text[:50]}...'")

    # Count paragraphs by style
    style_counts = {}
    for item in iter_block_items(doc):
        if isinstance(item, Paragraph):
            style = item.style.name if item.style else "None"
            style_counts[style] = style_counts.get(style, 0) + 1

    print("\nParagraph counts by style:")
    for style, count in sorted(style_counts.items()):
        print(f"  {style}: {count}")

    # Sample formatting check
    print("\nSample paragraph formatting verification:")
    para_idx = 0
    for item in iter_block_items(doc):
        if isinstance(item, Paragraph) and item.text.strip():
            if para_idx < 5 or "Investment" in item.text or "Dr. Sarah" in item.text:
                style = item.style.name if item.style else "None"
                print(f"\n  Para: '{item.text[:50]}...'")
                print(f"    Style: {style}")

                # Get style-based formatting
                if item.style:
                    try:
                        style_font = item.style.font
                        print(f"    Style font: name={style_font.name}, size={style_font.size.pt if style_font.size else None}pt, bold={style_font.bold}")
                    except:
                        pass

                # Check runs for color (like the red highlighted values)
                for r in item.runs:
                    if r.font.color and r.font.color.rgb:
                        print(f"    Run with color: '{r.text}' color={r.font.color.rgb}")
            para_idx += 1


def compare_with_original():
    """Compare the full sample with original template structure."""
    print(f"\n{'=' * 60}")
    print(" COMPARISON: FULL SAMPLE vs ORIGINAL TEMPLATE")
    print("=" * 60)

    orig = Document("test-proposal-template.docx")
    new = Document("full_sample_output.docx")

    # Get styles from both
    orig_styles = set()
    for item in iter_block_items(orig):
        if isinstance(item, Paragraph) and item.style:
            orig_styles.add(item.style.name)

    new_styles = set()
    for item in iter_block_items(new):
        if isinstance(item, Paragraph) and item.style:
            new_styles.add(item.style.name)

    print("\nStyles used in original template:", sorted(orig_styles))
    print("Styles used in full sample:", sorted(new_styles))
    print("Styles in both:", sorted(orig_styles & new_styles))
    print("Styles only in original:", sorted(orig_styles - new_styles))
    print("Styles only in sample:", sorted(new_styles - orig_styles))

    # Compare header/footer preservation
    print("\nHeader/Footer comparison:")
    print(f"  Original header: '{orig.sections[0].header.paragraphs[0].text}'")
    print(f"  Sample header:   '{new.sections[0].header.paragraphs[0].text}'")
    print(f"  Match: {orig.sections[0].header.paragraphs[0].text == new.sections[0].header.paragraphs[0].text}")


def main():
    print("=" * 60)
    print(" FULL TEMPLATE STYLE USAGE TEST")
    print("=" * 60)

    output_path = create_sample_proposal()
    analyze_output(output_path)
    compare_with_original()

    print("\n" + "=" * 60)
    print(" TEST COMPLETE")
    print("=" * 60)
    print("""
Files created:
  - full_sample_docxlang.json : DocxLang with all template styles used
  - full_sample_output.docx   : Generated document using all styles

The full sample document should:
1. Have the same header/footer as the template
2. Use all 6 paragraph styles (S1-S6) from the template
3. Include page breaks
4. Include colored text (red highlighted values)
5. Maintain the template's visual branding
""")


if __name__ == "__main__":
    main()
