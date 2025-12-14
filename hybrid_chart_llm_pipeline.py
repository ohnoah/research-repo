
"""
hybrid_chart_llm_pipeline.py

Hybrid "LLM-driven chart data update" pipeline for PPTX templates.

Key idea:
- For classic chart types that python-pptx can round-trip cleanly, use a semantic "DSL"
  and apply changes via python-pptx Chart.replace_data() (updates chart XML cache AND
  the embedded workbook).
- For everything else (including Office 2016+ cx: charts like treemap/sunburst/etc.),
  fall back to a workbook-cell patch approach using Aspose.Slides, which can edit the
  embedded chart workbook directly.

One LLM call per chart:
- You extract a per-chart prompt payload (includes chart type + current data).
- LLM returns an update JSON matching one of the Pydantic output models.
- You apply the update to a working copy (template slide cloned into output deck).

Requirements:
    pip install python-pptx openpyxl pydantic aspose-slides lxml
"""

from __future__ import annotations

import io
import json
import re
from dataclasses import dataclass
from enum import Enum
from typing import Any, Dict, Iterable, List, Literal, Optional, Sequence, Tuple, Union

import openpyxl
from openpyxl.utils.cell import (
    coordinate_from_string,
    column_index_from_string,
    get_column_letter,
    range_boundaries,
)

# --- Pydantic compat (v1/v2) ---
try:
    from pydantic.v1 import BaseModel, Field, ValidationError, root_validator, validator
except Exception:  # pragma: no cover
    from pydantic import BaseModel, Field, ValidationError, root_validator, validator

# lxml is used to parse chart XML
from lxml import etree as ET


# =============================================================================
# Strategy / Locators
# =============================================================================

class ChartStrategy(str, Enum):
    DSL = "dsl"
    WORKBOOK = "workbook"


class ChartLocator(BaseModel):
    """
    Identifies a chart shape on a slide.

    We use:
      - slide_index: 0-based index
      - shape_id: slide-scoped interop id (Aspose Shape.office_interop_shape_id; python-pptx shape.shape_id)
    """
    slide_index: int = Field(..., ge=0)
    shape_id: int = Field(..., ge=1)

    # Optional metadata to help debugging/prompting
    shape_name: Optional[str] = None
    alt_text: Optional[str] = None


# =============================================================================
# Prompt payload schemas (what you send to the LLM)
# =============================================================================

JsonScalar = Union[str, int, float, bool, None]


class DslSeriesCategory(BaseModel):
    name: str
    values: List[Optional[float]]


class CategoryChartDSL(BaseModel):
    kind: Literal["category_chart_dsl"] = "category_chart_dsl"
    chart_type: str
    categories: List[str]
    series: List[DslSeriesCategory]

    @root_validator
    def _validate_rectangular(cls, values: Dict[str, Any]) -> Dict[str, Any]:
        cats = values.get("categories") or []
        series = values.get("series") or []
        for s in series:
            if len(s.values) != len(cats):
                raise ValueError(
                    f"Series '{s.name}' has {len(s.values)} values but there are {len(cats)} categories."
                )
        return values


class DslSeriesXy(BaseModel):
    name: str
    points: List[Tuple[float, float]]  # JSON array [x, y]

    @validator("points")
    def _points_nonempty(cls, v: List[Tuple[float, float]]) -> List[Tuple[float, float]]:
        for p in v:
            if len(p) != 2:
                raise ValueError("Each XY point must be [x, y]")
        return v


class XyChartDSL(BaseModel):
    kind: Literal["xy_chart_dsl"] = "xy_chart_dsl"
    chart_type: str
    series: List[DslSeriesXy]


class DslSeriesBubble(BaseModel):
    name: str
    points: List[Tuple[float, float, float]]  # JSON array [x, y, size]

    @validator("points")
    def _points_shape(cls, v: List[Tuple[float, float, float]]) -> List[Tuple[float, float, float]]:
        for p in v:
            if len(p) != 3:
                raise ValueError("Each bubble point must be [x, y, size]")
        return v


class BubbleChartDSL(BaseModel):
    kind: Literal["bubble_chart_dsl"] = "bubble_chart_dsl"
    chart_type: str
    series: List[DslSeriesBubble]


DslPayload = Union[CategoryChartDSL, XyChartDSL, BubbleChartDSL]


class WorkbookGridSnapshot(BaseModel):
    """
    A rectangular slice of the embedded workbook for LLM I/O.

    NOTE: This is NOT a generic "Excel dump". It is deliberately bounded to
    the cells referenced by the chart formulas (when we can determine that).
    """
    sheet_name: str
    sheet_index: int = Field(..., ge=0)
    start_cell: str  # e.g., "A1"
    n_rows: int = Field(..., ge=1)
    n_cols: int = Field(..., ge=1)
    values: List[List[JsonScalar]]

    @validator("start_cell")
    def _valid_a1(cls, v: str) -> str:
        if not re.match(r"^[A-Z]{1,3}[1-9][0-9]*$", v):
            raise ValueError("start_cell must be A1 reference like 'A1'")
        return v

    @root_validator
    def _validate_rect(cls, values: Dict[str, Any]) -> Dict[str, Any]:
        grid = values.get("values") or []
        n_rows = values.get("n_rows")
        n_cols = values.get("n_cols")
        if len(grid) != n_rows:
            raise ValueError(f"values has {len(grid)} rows but n_rows={n_rows}")
        for row in grid:
            if len(row) != n_cols:
                raise ValueError("values must be rectangular with n_cols columns per row")
        return values


class WorkbookSnapshot(BaseModel):
    """
    Workbook-based prompt payload: send chart type + a small grid + editable cell list.
    """
    kind: Literal["workbook_snapshot"] = "workbook_snapshot"
    chart_type: str  # from Aspose chart.type (ChartType enum name) when available
    chart_data_source: Literal["embedded", "external", "unknown"] = "embedded"
    grid: WorkbookGridSnapshot
    editable_cells: List[str]  # A1 addresses safe to write (within chart data ranges, excluding formulas)


class ChartPromptPayload(BaseModel):
    """
    Unified per-chart LLM input payload.
    """
    locator: ChartLocator
    strategy: ChartStrategy
    chart_type: str

    # One of these is populated based on strategy
    dsl: Optional[DslPayload] = None
    workbook: Optional[WorkbookSnapshot] = None

    @root_validator
    def _validate_by_strategy(cls, values: Dict[str, Any]) -> Dict[str, Any]:
        strat = values.get("strategy")
        if strat == ChartStrategy.DSL and values.get("dsl") is None:
            raise ValueError("strategy=dsl requires dsl payload")
        if strat == ChartStrategy.WORKBOOK and values.get("workbook") is None:
            raise ValueError("strategy=workbook requires workbook payload")
        return values


# =============================================================================
# LLM output schemas (what LLM returns)
# =============================================================================

class CategoryChartUpdate(BaseModel):
    kind: Literal["category_chart_update"] = "category_chart_update"
    categories: List[str]
    series: List[DslSeriesCategory]

    @root_validator
    def _validate_rectangular(cls, values: Dict[str, Any]) -> Dict[str, Any]:
        cats = values.get("categories") or []
        series = values.get("series") or []
        for s in series:
            if len(s.values) != len(cats):
                raise ValueError(
                    f"Series '{s.name}' has {len(s.values)} values but there are {len(cats)} categories."
                )
        return values


class XyChartUpdate(BaseModel):
    kind: Literal["xy_chart_update"] = "xy_chart_update"
    series: List[DslSeriesXy]


class BubbleChartUpdate(BaseModel):
    kind: Literal["bubble_chart_update"] = "bubble_chart_update"
    series: List[DslSeriesBubble]


class CellPatch(BaseModel):
    address: str
    value: JsonScalar

    @validator("address")
    def _valid_a1(cls, v: str) -> str:
        if not re.match(r"^[A-Z]{1,3}[1-9][0-9]*$", v):
            raise ValueError("address must be an A1 reference like 'B2'")
        return v


class WorkbookCellPatchUpdate(BaseModel):
    kind: Literal["workbook_cell_patch_update"] = "workbook_cell_patch_update"
    sheet_index: int = Field(0, ge=0)
    cells: List[CellPatch]


class WorkbookGridReplaceUpdate(BaseModel):
    kind: Literal["workbook_grid_replace_update"] = "workbook_grid_replace_update"
    sheet_index: int = Field(0, ge=0)
    start_cell: str
    values: List[List[JsonScalar]]

    @validator("start_cell")
    def _valid_a1(cls, v: str) -> str:
        if not re.match(r"^[A-Z]{1,3}[1-9][0-9]*$", v):
            raise ValueError("start_cell must be an A1 reference like 'A1'")
        return v

    @root_validator
    def _validate_rect(cls, values: Dict[str, Any]) -> Dict[str, Any]:
        grid = values.get("values") or []
        if not grid:
            raise ValueError("values must not be empty")
        n_cols = len(grid[0])
        for row in grid:
            if len(row) != n_cols:
                raise ValueError("values must be rectangular")
        return values


ChartUpdate = Union[
    CategoryChartUpdate,
    XyChartUpdate,
    BubbleChartUpdate,
    WorkbookCellPatchUpdate,
    WorkbookGridReplaceUpdate,
]


# =============================================================================
# A1 helpers
# =============================================================================

def _a1_to_rowcol(a1: str) -> Tuple[int, int]:
    col_letters, row = coordinate_from_string(a1)
    return int(row), int(column_index_from_string(col_letters))


def _rowcol_to_a1(row: int, col: int) -> str:
    return f"{get_column_letter(col)}{row}"


def _json_safe(v: Any) -> JsonScalar:
    if v is None:
        return None
    if isinstance(v, (str, int, float, bool)):
        return v
    # datetime/date -> ISO string
    try:
        import datetime
        if isinstance(v, (datetime.date, datetime.datetime)):
            return v.isoformat()
    except Exception:
        pass
    return str(v)


# =============================================================================
# Slide cloning and chart discovery (Aspose)
# =============================================================================

def clone_slide_aspose(template_pptx: str, source_slide_index: int, out_pptx: str) -> int:
    """
    Clone a slide into a working output PPTX and return the new slide index.
    """
    import aspose.slides as slides  # type: ignore

    with slides.Presentation(str(template_pptx)) as pres:
        src_slide = pres.slides[source_slide_index]
        try:
            _ = pres.slides.add_clone(src_slide)  # type: ignore[attr-defined]
        except Exception:
            blank_layout = pres.masters[0].layout_slides.get_by_type(slides.SlideLayoutType.BLANK)
            new_slide = pres.slides.add_empty_slide(blank_layout)
            for shp in src_slide.shapes:
                new_slide.shapes.add_clone(shp)

        new_index = pres.slides.length - 1
        pres.save(str(out_pptx), slides.export.SaveFormat.PPTX)
    return int(new_index)


def list_chart_locators_aspose(pptx_path: str, slide_index: int) -> List[ChartLocator]:
    """
    Return a ChartLocator for every chart shape on the given slide.

    We use Aspose because python-pptx may not recognize "cx:" chart types.
    """
    import aspose.slides as slides  # type: ignore

    out: List[ChartLocator] = []
    with slides.Presentation(str(pptx_path)) as pres:
        slide = pres.slides[slide_index]
        for shape in slide.shapes:
            try:
                _ = shape.chart_data  # type: ignore[attr-defined]
            except Exception:
                continue

            try:
                sid = int(shape.office_interop_shape_id)  # type: ignore[attr-defined]
            except Exception:
                continue

            out.append(
                ChartLocator(
                    slide_index=slide_index,
                    shape_id=sid,
                    shape_name=getattr(shape, "name", None),
                    alt_text=getattr(shape, "alternative_text", None),
                )
            )
    return out


def get_chart_type_aspose(pptx_path: str, locator: ChartLocator) -> Tuple[str, Literal["embedded", "external", "unknown"]]:
    """
    Return (chart_type_name, data_source_kind) using Aspose.

    Aspose IChart has property .type (ChartType enum) and chart.chart_data.data_source_type.
    """
    import aspose.slides as slides  # type: ignore
    import aspose.slides.charts as charts  # type: ignore

    with slides.Presentation(str(pptx_path)) as pres:
        slide = pres.slides[locator.slide_index]
        chart_shape = None
        for shape in slide.shapes:
            try:
                if int(shape.office_interop_shape_id) != locator.shape_id:  # type: ignore[attr-defined]
                    continue
                _ = shape.chart_data  # type: ignore[attr-defined]
                chart_shape = shape
                break
            except Exception:
                continue

        if chart_shape is None:
            return ("UNKNOWN", "unknown")

        chart_type = "UNKNOWN"
        try:
            ct = chart_shape.type  # type: ignore[attr-defined]
            # Aspose ChartType is a .NET enum/flag; `str(ct)` often yields the numeric value.
            chart_type = getattr(ct, "name", None) or str(ct)
        except Exception:
            chart_type = "UNKNOWN"

        src_kind: Literal["embedded", "external", "unknown"] = "unknown"
        try:
            dst = chart_shape.chart_data.data_source_type  # type: ignore[attr-defined]
            if dst == charts.ChartDataSourceType.EXTERNAL_WORKBOOK:
                src_kind = "external"
            elif dst == charts.ChartDataSourceType.INTERNAL_WORKBOOK:
                src_kind = "embedded"
            else:
                src_kind = "unknown"
        except Exception:
            src_kind = "unknown"

        return (chart_type, src_kind)


# =============================================================================
# python-pptx chart access
# =============================================================================

def _get_chart_python_pptx(pptx_path: str, locator: ChartLocator):
    """
    Locate python-pptx chart object by slide index + shape_id.
    Returns (chart, shape, prs)
    """
    from pptx import Presentation  # type: ignore

    prs = Presentation(str(pptx_path))
    slide = prs.slides[locator.slide_index]
    for shape in slide.shapes:
        try:
            if shape.shape_id != locator.shape_id:
                continue
        except Exception:
            continue
        try:
            if shape.has_chart:
                return shape.chart, shape, prs
        except Exception:
            continue
    return None, None, prs


# =============================================================================
# Strategy detection
# =============================================================================

def choose_strategy(pptx_path: str, locator: ChartLocator) -> ChartStrategy:
    """
    Conservative decision: choose DSL only for chart types we can round-trip with python-pptx.

    DSL route supports:
      - category charts
      - XY scatter charts
      - bubble charts

    Everything else => workbook route (Aspose).
    """
    from pptx.enum.chart import XL_CHART_TYPE  # type: ignore

    CATEGORY_TYPES = {
        XL_CHART_TYPE.PIE,
        XL_CHART_TYPE.PIE_EXPLODED,
        XL_CHART_TYPE.PIE_OF_PIE,
        XL_CHART_TYPE.BAR_OF_PIE,
        XL_CHART_TYPE.DOUGHNUT,
        XL_CHART_TYPE.DOUGHNUT_EXPLODED,
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        XL_CHART_TYPE.COLUMN_STACKED,
        XL_CHART_TYPE.COLUMN_STACKED_100,
        XL_CHART_TYPE.BAR_CLUSTERED,
        XL_CHART_TYPE.BAR_STACKED,
        XL_CHART_TYPE.BAR_STACKED_100,
        XL_CHART_TYPE.LINE,
        XL_CHART_TYPE.LINE_MARKERS,
        XL_CHART_TYPE.LINE_STACKED,
        XL_CHART_TYPE.LINE_STACKED_100,
        XL_CHART_TYPE.AREA,
        XL_CHART_TYPE.AREA_STACKED,
        XL_CHART_TYPE.AREA_STACKED_100,
        XL_CHART_TYPE.RADAR,
        XL_CHART_TYPE.RADAR_FILLED,
        XL_CHART_TYPE.RADAR_MARKERS,
    }

    XY_TYPES = {
        XL_CHART_TYPE.XY_SCATTER,
        XL_CHART_TYPE.XY_SCATTER_LINES,
        XL_CHART_TYPE.XY_SCATTER_LINES_NO_MARKERS,
        XL_CHART_TYPE.XY_SCATTER_SMOOTH,
        XL_CHART_TYPE.XY_SCATTER_SMOOTH_NO_MARKERS,
    }

    BUBBLE_TYPES = {XL_CHART_TYPE.BUBBLE}
    # python-pptx enum naming differs across versions (e.g. 1.0.2 uses BUBBLE_THREE_D_EFFECT).
    bubble_3d = getattr(XL_CHART_TYPE, "BUBBLE_3D_EFFECT", None) or getattr(XL_CHART_TYPE, "BUBBLE_THREE_D_EFFECT", None)
    if bubble_3d is not None:
        BUBBLE_TYPES.add(bubble_3d)

    chart, _shape, _prs = _get_chart_python_pptx(pptx_path, locator)
    if chart is None:
        return ChartStrategy.WORKBOOK

    # Avoid multi-plot combo charts in DSL route (can be added later)
    try:
        if len(chart.plots) != 1:
            return ChartStrategy.WORKBOOK
    except Exception:
        return ChartStrategy.WORKBOOK

    try:
        ct = chart.chart_type
    except Exception:
        return ChartStrategy.WORKBOOK

    if ct in CATEGORY_TYPES or ct in XY_TYPES or ct in BUBBLE_TYPES:
        return ChartStrategy.DSL

    return ChartStrategy.WORKBOOK


# =============================================================================
# Chart XML parsing helpers (for XY/bubble extraction + workbook slicing)
# =============================================================================

CHART_NS = {"c": "http://schemas.openxmlformats.org/drawingml/2006/chart"}


def _get_chart_xml_root_from_python_pptx_chart(chart) -> ET._Element:
    """
    Returns parsed XML root element for the chart part.
    """
    xml_bytes = None
    try:
        xml_bytes = chart.part.blob  # type: ignore[attr-defined]
    except Exception:
        pass

    if not xml_bytes:
        try:
            xml_bytes = ET.tostring(chart.part._element, encoding="utf-8")  # type: ignore[attr-defined]
        except Exception as e:
            raise RuntimeError("Unable to obtain chart XML from python-pptx chart") from e

    return ET.fromstring(xml_bytes)


def _pt_cache_to_list(pt_elements: Sequence[ET._Element]) -> List[Optional[float]]:
    """
    Given list of <c:pt idx="..."><c:v>...</c:v></c:pt>, return ordered list by idx.
    """
    idx_map: Dict[int, Optional[float]] = {}
    for pt in pt_elements:
        idx_s = pt.get("idx")
        if idx_s is None:
            continue
        try:
            idx = int(idx_s)
        except Exception:
            continue
        v = pt.findtext("c:v", namespaces=CHART_NS)
        if v is None or v == "":
            idx_map[idx] = None
        else:
            try:
                idx_map[idx] = float(v)
            except Exception:
                idx_map[idx] = None
    return [idx_map[i] for i in sorted(idx_map.keys())]

def _pt_cache_to_str_list(pt_elements: Sequence[ET._Element]) -> List[str]:
    """
    Given list of <c:pt idx="..."><c:v>...</c:v></c:pt>, return ordered list of strings by idx.
    Missing values become "".
    """
    idx_map: Dict[int, str] = {}
    for pt in pt_elements:
        idx_s = pt.get("idx")
        if idx_s is None:
            continue
        try:
            idx = int(idx_s)
        except Exception:
            continue
        v = pt.findtext("c:v", namespaces=CHART_NS)
        idx_map[idx] = "" if v is None else str(v)
    return [idx_map[i] for i in sorted(idx_map.keys())]


def _extract_series_name_from_ser(ser: ET._Element) -> str:
    """
    Best-effort series name extraction from a <c:ser> node.
    """
    name = ser.xpath("./c:tx/c:v/text()", namespaces=CHART_NS)
    if name and name[0].strip():
        return name[0].strip()

    cached = ser.xpath(".//c:tx//c:strCache//c:pt[@idx='0']/c:v/text()", namespaces=CHART_NS)
    if cached and cached[0].strip():
        return cached[0].strip()

    return ""


def extract_xy_chart_dsl_from_chart_xml(pptx_path: str, locator: ChartLocator) -> XyChartDSL:
    """
    Extract XY (scatter) chart data from chart XML caches.
    """
    chart, _shape, _prs = _get_chart_python_pptx(pptx_path, locator)
    if chart is None:
        raise RuntimeError("Chart not found via python-pptx; cannot extract XY DSL.")

    root = _get_chart_xml_root_from_python_pptx_chart(chart)

    chart_type = getattr(chart.chart_type, "name", str(chart.chart_type))
    series_nodes = root.xpath(".//c:ser", namespaces=CHART_NS)

    series_payload: List[DslSeriesXy] = []
    for ser in series_nodes:
        name = _extract_series_name_from_ser(ser)

        x_pts = ser.xpath("./c:xVal//c:numCache/c:pt", namespaces=CHART_NS)
        y_pts = ser.xpath("./c:yVal//c:numCache/c:pt", namespaces=CHART_NS)
        xs = _pt_cache_to_list(x_pts)
        ys = _pt_cache_to_list(y_pts)

        n = min(len(xs), len(ys))
        points: List[Tuple[float, float]] = []
        for i in range(n):
            x = xs[i]
            y = ys[i]
            if x is None or y is None:
                continue
            points.append((float(x), float(y)))

        series_payload.append(DslSeriesXy(name=name or f"Series {len(series_payload)+1}", points=points))

    return XyChartDSL(chart_type=chart_type, series=series_payload)


def extract_bubble_chart_dsl_from_chart_xml(pptx_path: str, locator: ChartLocator) -> BubbleChartDSL:
    """
    Extract bubble chart data from chart XML caches (xVal, yVal, bubbleSize).
    """
    chart, _shape, _prs = _get_chart_python_pptx(pptx_path, locator)
    if chart is None:
        raise RuntimeError("Chart not found via python-pptx; cannot extract Bubble DSL.")

    root = _get_chart_xml_root_from_python_pptx_chart(chart)

    chart_type = getattr(chart.chart_type, "name", str(chart.chart_type))
    series_nodes = root.xpath(".//c:ser", namespaces=CHART_NS)

    series_payload: List[DslSeriesBubble] = []
    for ser in series_nodes:
        name = _extract_series_name_from_ser(ser)

        x_pts = ser.xpath("./c:xVal//c:numCache/c:pt", namespaces=CHART_NS)
        y_pts = ser.xpath("./c:yVal//c:numCache/c:pt", namespaces=CHART_NS)
        s_pts = ser.xpath("./c:bubbleSize//c:numCache/c:pt", namespaces=CHART_NS)

        xs = _pt_cache_to_list(x_pts)
        ys = _pt_cache_to_list(y_pts)
        ss = _pt_cache_to_list(s_pts)

        n = min(len(xs), len(ys), len(ss))
        points: List[Tuple[float, float, float]] = []
        for i in range(n):
            x, y, sz = xs[i], ys[i], ss[i]
            if x is None or y is None or sz is None:
                continue
            points.append((float(x), float(y), float(sz)))

        series_payload.append(DslSeriesBubble(name=name or f"Series {len(series_payload)+1}", points=points))

    return BubbleChartDSL(chart_type=chart_type, series=series_payload)

def extract_category_chart_dsl_from_chart_xml(pptx_path: str, locator: ChartLocator) -> CategoryChartDSL:
    """
    Extract category chart data from chart XML caches.

    This is a fallback for templates where python-pptx cannot reliably surface categories
    via its public API (e.g., some charts backed by XLSB workbooks).
    """
    chart, _shape, _prs = _get_chart_python_pptx(pptx_path, locator)
    if chart is None:
        raise RuntimeError("Chart not found via python-pptx; cannot extract Category DSL.")

    root = _get_chart_xml_root_from_python_pptx_chart(chart)
    chart_type = getattr(chart.chart_type, "name", str(chart.chart_type))

    series_nodes = root.xpath(".//c:ser", namespaces=CHART_NS)

    # Categories: look for string or numeric caches under the first series.
    categories: List[str] = []
    if series_nodes:
        cat_str_pts = series_nodes[0].xpath("./c:cat//c:strCache/c:pt", namespaces=CHART_NS)
        if cat_str_pts:
            categories = [s for s in _pt_cache_to_str_list(cat_str_pts)]
        else:
            cat_num_pts = series_nodes[0].xpath("./c:cat//c:numCache/c:pt", namespaces=CHART_NS)
            if cat_num_pts:
                categories = [s for s in _pt_cache_to_str_list(cat_num_pts)]

        if not categories:
            # multi-level categories sometimes store as multiLvlStrCache
            ml_pts = series_nodes[0].xpath("./c:cat//c:multiLvlStrCache//c:pt", namespaces=CHART_NS)
            if ml_pts:
                categories = [s for s in _pt_cache_to_str_list(ml_pts)]

    # Series values.
    series_payload: List[DslSeriesCategory] = []
    max_len = 0
    raw_values: List[Tuple[str, List[Optional[float]]]] = []
    for ser in series_nodes:
        name = _extract_series_name_from_ser(ser)
        v_pts = ser.xpath("./c:val//c:numCache/c:pt", namespaces=CHART_NS)
        vals = _pt_cache_to_list(v_pts)
        max_len = max(max_len, len(vals))
        raw_values.append((name or f"Series {len(raw_values)+1}", vals))

    if not categories:
        categories = [f"Category {i+1}" for i in range(max_len)]

    # Normalize all series to match category length (pad/trim).
    for name, vals in raw_values:
        if len(vals) < len(categories):
            vals = vals + [None] * (len(categories) - len(vals))
        elif len(vals) > len(categories):
            vals = vals[: len(categories)]
        series_payload.append(DslSeriesCategory(name=name, values=vals))

    return CategoryChartDSL(chart_type=chart_type, categories=categories, series=series_payload)


def extract_category_chart_dsl_python_pptx(pptx_path: str, locator: ChartLocator) -> CategoryChartDSL:
    """
    Extract categories + series values for category charts using python-pptx public API.
    """
    chart, _shape, _prs = _get_chart_python_pptx(pptx_path, locator)
    if chart is None:
        raise RuntimeError("Chart not found via python-pptx; cannot extract Category DSL.")

    try:
        plot = chart.plots[0]
        categories = [str(cat) for cat in list(plot.categories)]

        series_payload: List[DslSeriesCategory] = []
        for s in chart.series:
            vals: List[Optional[float]] = []
            for v in s.values:
                if v is None:
                    vals.append(None)
                else:
                    vals.append(float(v))
            series_payload.append(DslSeriesCategory(name=s.name, values=vals))

        chart_type = getattr(chart.chart_type, "name", str(chart.chart_type))
        return CategoryChartDSL(chart_type=chart_type, categories=categories, series=series_payload)
    except Exception:
        # Fall back to XML-cache extraction when python-pptx can't surface categories.
        return extract_category_chart_dsl_from_chart_xml(pptx_path, locator)


# =============================================================================
# Chart-to-workbook "range slicer" (chart-type-specific workbook snapshot)
# =============================================================================

@dataclass(frozen=True)
class _ChartRangeRef:
    sheet_name: str
    a1_range: str
    kind: Literal["num", "str"]


def _normalize_sheet_name(sheet: str) -> str:
    s = sheet.strip()
    s = re.sub(r"^\[[0-9]+\]", "", s).strip()
    if len(s) >= 2 and s[0] == "'" and s[-1] == "'":
        s = s[1:-1]
    return s


def _parse_chart_formula_ref(formula: str) -> Optional[Tuple[str, str]]:
    """
    Parse a chart formula reference like:
      Sheet1!$A$2:$B$5
      'My Sheet'!$A$2
    Returns (sheet_name, a1_range_without_dollars)
    """
    if not formula or "!" not in formula:
        return None
    sheet_part, ref_part = formula.split("!", 1)
    sheet_name = _normalize_sheet_name(sheet_part)
    ref_part = ref_part.split(",", 1)[0]
    ref = ref_part.replace("$", "").strip()
    if not re.match(r"^[A-Z]{1,3}[1-9][0-9]*(?::[A-Z]{1,3}[1-9][0-9]*)?$", ref):
        return None
    return sheet_name, ref


def extract_chart_range_refs_from_python_pptx(pptx_path: str, locator: ChartLocator) -> List[_ChartRangeRef]:
    """
    Extract all worksheet range references used by the chart from chart XML.
    """
    chart, _shape, _prs = _get_chart_python_pptx(pptx_path, locator)
    if chart is None:
        return []

    try:
        root = _get_chart_xml_root_from_python_pptx_chart(chart)
    except Exception:
        return []

    refs: List[_ChartRangeRef] = []

    for f_el in root.xpath(".//c:numRef/c:f", namespaces=CHART_NS):
        f_txt = (f_el.text or "").strip()
        parsed = _parse_chart_formula_ref(f_txt)
        if parsed:
            sheet, a1 = parsed
            refs.append(_ChartRangeRef(sheet_name=sheet, a1_range=a1, kind="num"))

    for f_el in root.xpath(".//c:strRef/c:f | .//c:multiLvlStrRef/c:f", namespaces=CHART_NS):
        f_txt = (f_el.text or "").strip()
        parsed = _parse_chart_formula_ref(f_txt)
        if parsed:
            sheet, a1 = parsed
            refs.append(_ChartRangeRef(sheet_name=sheet, a1_range=a1, kind="str"))

    return refs


def _pick_sheet_for_refs(workbook: openpyxl.Workbook, refs: List[_ChartRangeRef], fallback_index: int = 0) -> Tuple[int, str]:
    if refs:
        counts: Dict[str, int] = {}
        for r in refs:
            counts[r.sheet_name] = counts.get(r.sheet_name, 0) + 1
        for sheet_name, _count in sorted(counts.items(), key=lambda kv: kv[1], reverse=True):
            if sheet_name in workbook.sheetnames:
                idx = workbook.sheetnames.index(sheet_name)
                return idx, sheet_name

    fallback_index = min(max(fallback_index, 0), len(workbook.worksheets) - 1)
    ws = workbook.worksheets[fallback_index]
    return fallback_index, ws.title


def _bounding_box_from_refs(refs: List[_ChartRangeRef], sheet_name: str) -> Optional[Tuple[int, int, int, int]]:
    sheet_refs = [r for r in refs if r.sheet_name == sheet_name]
    if not sheet_refs:
        return None

    min_row, min_col = 10**9, 10**9
    max_row, max_col = 0, 0

    for r in sheet_refs:
        try:
            min_c, min_r, max_c, max_r = range_boundaries(r.a1_range)
        except Exception:
            continue
        min_row = min(min_row, min_r)
        min_col = min(min_col, min_c)
        max_row = max(max_row, max_r)
        max_col = max(max_col, max_c)

    if max_row == 0:
        return None
    return (min_row, min_col, max_row, max_col)


def _clamp_box(box: Tuple[int, int, int, int], max_rows: int, max_cols: int) -> Tuple[int, int, int, int]:
    min_row, min_col, max_row, max_col = box
    min_row = max(1, min_row)
    min_col = max(1, min_col)
    max_row = min(max_row, max_rows)
    max_col = min(max_col, max_cols)
    return min_row, min_col, max_row, max_col


def _expand_box(
    box: Tuple[int, int, int, int],
    pad_rows_top: int = 1,
    pad_cols_left: int = 1,
    pad_rows_bottom: int = 0,
    pad_cols_right: int = 0,
) -> Tuple[int, int, int, int]:
    min_row, min_col, max_row, max_col = box
    return (
        max(1, min_row - pad_rows_top),
        max(1, min_col - pad_cols_left),
        max_row + pad_rows_bottom,
        max_col + pad_cols_right,
    )


def _iter_cells_in_range(a1_range: str) -> Iterable[str]:
    min_c, min_r, max_c, max_r = range_boundaries(a1_range)
    for r in range(min_r, max_r + 1):
        for c in range(min_c, max_c + 1):
            yield _rowcol_to_a1(r, c)


def read_embedded_workbook_bytes_aspose(pptx_path: str, locator: ChartLocator) -> bytes:
    """
    Read the embedded chart workbook bytes via Aspose.
    """
    import aspose.slides as slides  # type: ignore
    with slides.Presentation(str(pptx_path)) as pres:
        slide = pres.slides[locator.slide_index]
        chart_shape = None
        for shape in slide.shapes:
            try:
                if int(shape.office_interop_shape_id) != locator.shape_id:  # type: ignore[attr-defined]
                    continue
                _ = shape.chart_data  # type: ignore[attr-defined]
                chart_shape = shape
                break
            except Exception:
                continue

        if chart_shape is None:
            raise RuntimeError("Chart not found via Aspose; cannot read workbook stream.")

        return chart_shape.chart_data.read_workbook_stream().read()  # type: ignore[attr-defined]


def extract_workbook_snapshot(
    pptx_path: str,
    locator: ChartLocator,
    *,
    preferred_sheet_index: int = 0,
    max_rows: int = 120,
    max_cols: int = 30,
    prefer_chart_ranges: bool = True,
) -> WorkbookSnapshot:
    """
    Extract a workbook snapshot for the LLM.
    """
    chart_type, data_source_kind = get_chart_type_aspose(pptx_path, locator)

    wb_bytes = read_embedded_workbook_bytes_aspose(pptx_path, locator)
    wb = openpyxl.load_workbook(io.BytesIO(wb_bytes), data_only=False)

    refs: List[_ChartRangeRef] = []
    if prefer_chart_ranges:
        refs = extract_chart_range_refs_from_python_pptx(pptx_path, locator)

    sheet_index, sheet_name = _pick_sheet_for_refs(wb, refs, fallback_index=preferred_sheet_index)
    ws = wb.worksheets[sheet_index]

    box = _bounding_box_from_refs(refs, sheet_name) if refs else None

    if box:
        box = _expand_box(box, pad_rows_top=1, pad_cols_left=1)
        box = _clamp_box(box, max_rows=max_rows, max_cols=max_cols)
        min_row, min_col, max_row, max_col = box
    else:
        max_r = min(ws.max_row or 1, max_rows)
        max_c = min(ws.max_column or 1, max_cols)
        min_row, min_col = 1, 1
        max_row, max_col = max_r, max_c

    n_rows = max_row - min_row + 1
    n_cols = max_col - min_col + 1
    start_cell = f"{get_column_letter(min_col)}{min_row}"

    values: List[List[JsonScalar]] = []
    for r in range(min_row, max_row + 1):
        row_vals: List[JsonScalar] = []
        for c in range(min_col, max_col + 1):
            v = ws.cell(r, c).value
            row_vals.append(_json_safe(v))
        values.append(row_vals)

    editable: List[str] = []
    chosen_refs = [r for r in refs if r.sheet_name == sheet_name]
    for r in chosen_refs:
        for addr in _iter_cells_in_range(r.a1_range):
            cell = ws[addr]
            if isinstance(cell.value, str) and cell.value.startswith("="):
                continue
            editable.append(addr)

    if not editable:
        for r in range(min_row, max_row + 1):
            for c in range(min_col, max_col + 1):
                addr = _rowcol_to_a1(r, c)
                v = ws.cell(r, c).value
                if isinstance(v, (int, float)):
                    editable.append(addr)

    grid = WorkbookGridSnapshot(
        sheet_name=ws.title,
        sheet_index=sheet_index,
        start_cell=start_cell,
        n_rows=n_rows,
        n_cols=n_cols,
        values=values,
    )

    return WorkbookSnapshot(
        chart_type=chart_type,
        chart_data_source=data_source_kind,
        grid=grid,
        editable_cells=sorted(set(editable)),
    )


# =============================================================================
# Build prompt payload (strategy + extraction)
# =============================================================================

def build_llm_prompt_payload(pptx_path: str, locator: ChartLocator) -> ChartPromptPayload:
    """
    Main per-chart entrypoint:
      - decide DSL vs workbook route
      - extract current chart data (DSL) OR workbook snapshot
      - always include chart_type for prompt context
    """
    aspose_chart_type, _src_kind = get_chart_type_aspose(pptx_path, locator)

    strategy = choose_strategy(pptx_path, locator)
    if strategy == ChartStrategy.DSL:
        chart, _shape, _prs = _get_chart_python_pptx(pptx_path, locator)
        if chart is None:
            wb = extract_workbook_snapshot(pptx_path, locator)
            return ChartPromptPayload(locator=locator, strategy=ChartStrategy.WORKBOOK, chart_type=aspose_chart_type, workbook=wb)

        ct_name = getattr(chart.chart_type, "name", str(chart.chart_type))

        try:
            if ct_name.startswith("XY_SCATTER"):
                dsl = extract_xy_chart_dsl_from_chart_xml(pptx_path, locator)
                return ChartPromptPayload(locator=locator, strategy=strategy, chart_type=ct_name, dsl=dsl)
            if ct_name.startswith("BUBBLE"):
                dsl = extract_bubble_chart_dsl_from_chart_xml(pptx_path, locator)
                return ChartPromptPayload(locator=locator, strategy=strategy, chart_type=ct_name, dsl=dsl)

            dsl = extract_category_chart_dsl_python_pptx(pptx_path, locator)
            return ChartPromptPayload(locator=locator, strategy=strategy, chart_type=ct_name, dsl=dsl)

        except Exception:
            wb = extract_workbook_snapshot(pptx_path, locator)
            return ChartPromptPayload(locator=locator, strategy=ChartStrategy.WORKBOOK, chart_type=aspose_chart_type, workbook=wb)

    wb = extract_workbook_snapshot(pptx_path, locator)
    return ChartPromptPayload(locator=locator, strategy=strategy, chart_type=aspose_chart_type, workbook=wb)


# =============================================================================
# Apply updates back to PPTX
# =============================================================================

def apply_category_chart_update_python_pptx(
    pptx_in: str,
    pptx_out: str,
    locator: ChartLocator,
    update: CategoryChartUpdate,
) -> None:
    from pptx import Presentation  # type: ignore
    from pptx.chart.data import CategoryChartData  # type: ignore

    chart_data = CategoryChartData()
    chart_data.categories = update.categories
    for s in update.series:
        chart_data.add_series(s.name, s.values)

    def _apply(pptx_path: str, out_path: str) -> None:
        prs = Presentation(str(pptx_path))
        slide = prs.slides[locator.slide_index]

        target_shape = None
        for shape in slide.shapes:
            try:
                if shape.shape_id != locator.shape_id:
                    continue
                if not shape.has_chart:
                    continue
                target_shape = shape
                break
            except Exception:
                continue

        if target_shape is None:
            raise RuntimeError("Chart shape not found for category DSL update.")

        chart = target_shape.chart
        chart.replace_data(chart_data)
        prs.save(str(out_path))

    try:
        _apply(pptx_in, pptx_out)
    except ValueError as e:
        if "target-mode is external" in str(e):
            raise RuntimeError(
                "Chart workbook is EXTERNAL; run an offline internalization step before using DSL updates."
            ) from e
        raise


def apply_xy_chart_update_python_pptx(
    pptx_in: str,
    pptx_out: str,
    locator: ChartLocator,
    update: XyChartUpdate,
) -> None:
    from pptx import Presentation  # type: ignore
    from pptx.chart.data import XyChartData  # type: ignore

    chart_data = XyChartData()
    for s in update.series:
        sdata = chart_data.add_series(s.name)
        for x, y in s.points:
            sdata.add_data_point(x, y)

    def _apply(pptx_path: str, out_path: str) -> None:
        prs = Presentation(str(pptx_path))
        slide = prs.slides[locator.slide_index]

        target_shape = None
        for shape in slide.shapes:
            try:
                if shape.shape_id != locator.shape_id:
                    continue
                if not shape.has_chart:
                    continue
                target_shape = shape
                break
            except Exception:
                continue

        if target_shape is None:
            raise RuntimeError("Chart shape not found for XY DSL update.")

        chart = target_shape.chart
        chart.replace_data(chart_data)
        prs.save(str(out_path))

    try:
        _apply(pptx_in, pptx_out)
    except ValueError as e:
        if "target-mode is external" in str(e):
            raise RuntimeError(
                "Chart workbook is EXTERNAL; run an offline internalization step before using DSL updates."
            ) from e
        raise


def apply_bubble_chart_update_python_pptx(
    pptx_in: str,
    pptx_out: str,
    locator: ChartLocator,
    update: BubbleChartUpdate,
) -> None:
    from pptx import Presentation  # type: ignore
    from pptx.chart.data import BubbleChartData  # type: ignore

    chart_data = BubbleChartData()
    for s in update.series:
        sdata = chart_data.add_series(s.name)
        for x, y, size in s.points:
            sdata.add_data_point(x, y, size)

    def _apply(pptx_path: str, out_path: str) -> None:
        prs = Presentation(str(pptx_path))
        slide = prs.slides[locator.slide_index]

        target_shape = None
        for shape in slide.shapes:
            try:
                if shape.shape_id != locator.shape_id:
                    continue
                if not shape.has_chart:
                    continue
                target_shape = shape
                break
            except Exception:
                continue

        if target_shape is None:
            raise RuntimeError("Chart shape not found for Bubble DSL update.")

        chart = target_shape.chart
        chart.replace_data(chart_data)
        prs.save(str(out_path))

    try:
        _apply(pptx_in, pptx_out)
    except ValueError as e:
        if "target-mode is external" in str(e):
            raise RuntimeError(
                "Chart workbook is EXTERNAL; run an offline internalization step before using DSL updates."
            ) from e
        raise


def apply_workbook_update_aspose(
    pptx_in: str,
    pptx_out: str,
    locator: ChartLocator,
    update: Union[WorkbookCellPatchUpdate, WorkbookGridReplaceUpdate],
    *,
    enforce_editable_cells: Optional[List[str]] = None,
) -> None:
    import aspose.slides as slides  # type: ignore

    allowed = set(enforce_editable_cells or [])

    def _apply(path_in: str, path_out: str) -> None:
        with slides.Presentation(str(path_in)) as pres:
            slide = pres.slides[locator.slide_index]

            chart_shape = None
            for shape in slide.shapes:
                try:
                    if int(shape.office_interop_shape_id) != locator.shape_id:  # type: ignore[attr-defined]
                        continue
                    _ = shape.chart_data  # type: ignore[attr-defined]
                    chart_shape = shape
                    break
                except Exception:
                    continue

            if chart_shape is None:
                raise RuntimeError("Chart not found via Aspose; cannot apply workbook update.")

            workbook = chart_shape.chart_data.chart_data_workbook  # type: ignore[attr-defined]

            if isinstance(update, WorkbookCellPatchUpdate):
                for cp in update.cells:
                    if allowed and cp.address not in allowed:
                        raise ValueError(f"Refusing to write non-editable cell {cp.address}")
                    workbook.get_cell(update.sheet_index, cp.address, cp.value)

            elif isinstance(update, WorkbookGridReplaceUpdate):
                start_row, start_col = _a1_to_rowcol(update.start_cell)
                for r_off, row_vals in enumerate(update.values):
                    for c_off, v in enumerate(row_vals):
                        addr = _rowcol_to_a1(start_row + r_off, start_col + c_off)
                        if allowed and addr not in allowed:
                            raise ValueError(f"Refusing to write non-editable cell {addr}")
                        workbook.get_cell(update.sheet_index, addr, v)

            else:
                raise TypeError("Unsupported workbook update type")

            pres.save(str(path_out), slides.export.SaveFormat.PPTX)

    _apply(pptx_in, pptx_out)


def apply_chart_update(
    pptx_in: str,
    pptx_out: str,
    prompt_payload: ChartPromptPayload,
    llm_update_json: str,
) -> None:
    payload = json.loads(llm_update_json)
    kind = payload.get("kind")

    if kind == "category_chart_update":
        update = CategoryChartUpdate.parse_obj(payload)
        apply_category_chart_update_python_pptx(pptx_in, pptx_out, prompt_payload.locator, update)
        return

    if kind == "xy_chart_update":
        update = XyChartUpdate.parse_obj(payload)
        apply_xy_chart_update_python_pptx(pptx_in, pptx_out, prompt_payload.locator, update)
        return

    if kind == "bubble_chart_update":
        update = BubbleChartUpdate.parse_obj(payload)
        apply_bubble_chart_update_python_pptx(pptx_in, pptx_out, prompt_payload.locator, update)
        return

    if kind == "workbook_cell_patch_update":
        update = WorkbookCellPatchUpdate.parse_obj(payload)
        enforce = prompt_payload.workbook.editable_cells if prompt_payload.workbook else None
        apply_workbook_update_aspose(pptx_in, pptx_out, prompt_payload.locator, update, enforce_editable_cells=enforce)
        return

    if kind == "workbook_grid_replace_update":
        update = WorkbookGridReplaceUpdate.parse_obj(payload)
        enforce = prompt_payload.workbook.editable_cells if prompt_payload.workbook else None
        apply_workbook_update_aspose(pptx_in, pptx_out, prompt_payload.locator, update, enforce_editable_cells=enforce)
        return

    raise ValueError(f"Unknown update kind: {kind}")


# =============================================================================
# Example driver (no LLM call; prints payloads)
# =============================================================================

def example_extract_payloads_for_slide(template_pptx: str, template_slide_index: int, working_pptx: str) -> None:
    """
    Creates a working PPTX by cloning template_slide_index, then prints per-chart prompt payloads.
    """
    new_slide_idx = clone_slide_aspose(template_pptx, template_slide_index, working_pptx)
    charts = list_chart_locators_aspose(working_pptx, new_slide_idx)

    for loc in charts:
        payload = build_llm_prompt_payload(working_pptx, loc)
        print("\n=== CHART PAYLOAD ===")
        print(payload.json(indent=2))


if __name__ == "__main__":
    # example_extract_payloads_for_slide("template.pptx", 0, "working.pptx")
    pass
