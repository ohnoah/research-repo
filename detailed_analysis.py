#!/usr/bin/env python3
"""
Detailed analysis of round-trip differences and style issues.
"""

import json
from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph


def iter_block_items(doc: Document):
    """Yield Paragraph and Table objects in document order."""
    body = doc.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, doc)
        elif child.tag == qn("w:tbl"):
            yield Table(child, doc)


def analyze_paragraph_details(doc_path: str, label: str):
    """Get detailed paragraph info."""
    doc = Document(doc_path)
    paras = []

    for item in iter_block_items(doc):
        if isinstance(item, Paragraph):
            p = item
            info = {
                "text": (p.text or "")[:60],
                "style": p.style.name if p.style else None,
                "runs": len(p.runs),
                "run_details": [],
            }

            for r in p.runs[:5]:  # First 5 runs
                run_info = {
                    "text": (r.text or "")[:30],
                    "bold": r.bold,
                    "italic": r.italic,
                    "font_size": r.font.size.pt if r.font.size else None,
                    "font_name": r.font.name,
                }
                try:
                    run_info["style"] = r.style.name if r.style else None
                except:
                    run_info["style"] = None
                try:
                    run_info["color"] = str(r.font.color.rgb) if r.font.color and r.font.color.rgb else None
                except:
                    run_info["color"] = None
                info["run_details"].append(run_info)

            paras.append(info)

    return paras


def find_missing_style():
    """Find why Sub-header style is missing."""
    print("\n" + "=" * 60)
    print(" INVESTIGATING MISSING 'Sub-header (CVHP)' STYLE")
    print("=" * 60)

    orig_doc = Document("test-proposal-template.docx")

    print("\nLooking for 'Sub-header (CVHP)' in original document:")
    idx = 0
    for item in iter_block_items(orig_doc):
        if isinstance(item, Paragraph):
            style_name = item.style.name if item.style else None
            if style_name and "sub-header" in style_name.lower():
                print(f"  Found at para index {idx}: style='{style_name}', text='{(item.text or '')[:50]}'")
            idx += 1

    # Check if style exists in document styles
    print("\nStyles defined in document:")
    for style in orig_doc.styles:
        if "header" in style.name.lower() or "sub" in style.name.lower():
            print(f"  {style.name} (type: {style.type})")


def compare_specific_paragraphs():
    """Compare paragraphs one by one."""
    print("\n" + "=" * 60)
    print(" PARAGRAPH-BY-PARAGRAPH COMPARISON")
    print("=" * 60)

    orig_paras = analyze_paragraph_details("test-proposal-template.docx", "Original")
    rt_paras = analyze_paragraph_details("roundtrip_output.docx", "Roundtrip")

    print(f"\nOriginal: {len(orig_paras)} paragraphs")
    print(f"Roundtrip: {len(rt_paras)} paragraphs")

    # Find style differences
    print("\nStyle differences:")
    for i in range(min(len(orig_paras), len(rt_paras))):
        orig_style = orig_paras[i]["style"]
        rt_style = rt_paras[i]["style"]
        if orig_style != rt_style:
            print(f"  Para {i}: '{orig_style}' -> '{rt_style}'")
            print(f"    Text: '{orig_paras[i]['text'][:50]}'")

    # Find run count differences with details
    print("\nRun count differences (with formatting impact):")
    diff_count = 0
    for i in range(min(len(orig_paras), len(rt_paras))):
        orig_runs = orig_paras[i]["runs"]
        rt_runs = rt_paras[i]["runs"]
        if orig_runs != rt_runs:
            diff_count += 1
            if diff_count <= 10:
                print(f"  Para {i}: {orig_runs} runs -> {rt_runs} runs")
                print(f"    Text: '{orig_paras[i]['text'][:40]}'")
                # Show original run details
                if orig_paras[i]["run_details"]:
                    print(f"    Original runs:")
                    for j, rd in enumerate(orig_paras[i]["run_details"][:3]):
                        print(f"      Run {j}: '{rd['text'][:20]}' bold={rd['bold']} color={rd['color']}")
                if rt_paras[i]["run_details"]:
                    print(f"    Roundtrip runs:")
                    for j, rd in enumerate(rt_paras[i]["run_details"][:3]):
                        print(f"      Run {j}: '{rd['text'][:20]}' bold={rd['bold']} color={rd['color']}")

    if diff_count > 10:
        print(f"  ... and {diff_count - 10} more paragraphs with run differences")


def analyze_new_content_styling():
    """Analyze if new content document has proper styling."""
    print("\n" + "=" * 60)
    print(" NEW CONTENT STYLING ANALYSIS")
    print("=" * 60)

    new_doc = Document("new_content_output.docx")

    print("\nParagraphs in new content document:")
    for i, item in enumerate(iter_block_items(new_doc)):
        if isinstance(item, Paragraph):
            p = item
            style_name = p.style.name if p.style else None

            # Get actual formatting from paragraph
            pf = p.paragraph_format
            space_before = pf.space_before.pt if pf.space_before else None
            space_after = pf.space_after.pt if pf.space_after else None

            print(f"\n  Para {i}: style='{style_name}'")
            print(f"    Text: '{(p.text or '')[:50]}'")
            print(f"    Space before: {space_before}, Space after: {space_after}")

            for j, r in enumerate(p.runs[:3]):
                font = r.font
                print(f"    Run {j}: '{(r.text or '')[:20]}' font={font.name} size={font.size.pt if font.size else None} bold={r.bold}")

    # Check headers/footers preserved
    print("\n\nHeaders/Footers check:")
    for i, section in enumerate(new_doc.sections):
        header_text = " | ".join([p.text for p in section.header.paragraphs if p.text.strip()])
        footer_text = " | ".join([p.text for p in section.footer.paragraphs if p.text.strip()])
        print(f"  Section {i}:")
        print(f"    Header: '{header_text[:60]}'" if header_text else "    Header: (empty)")
        print(f"    Footer: '{footer_text[:60]}'" if footer_text else "    Footer: (empty)")


def check_page_break_handling():
    """Check how page breaks are handled."""
    print("\n" + "=" * 60)
    print(" PAGE BREAK HANDLING")
    print("=" * 60)

    with open("template_docxlang.json") as f:
        docxlang = json.load(f)

    print("\nPage breaks in DocxLang representation:")
    for i, block in enumerate(docxlang["blocks"]):
        if block.get("type") == "page_break":
            print(f"  Block {i}: page_break")
            if i > 0:
                print(f"    Previous block: {docxlang['blocks'][i-1].get('type')}")
            if i < len(docxlang["blocks"]) - 1:
                print(f"    Next block: {docxlang['blocks'][i+1].get('type')}")


def main():
    find_missing_style()
    compare_specific_paragraphs()
    analyze_new_content_styling()
    check_page_break_handling()

    print("\n" + "=" * 60)
    print(" SUMMARY OF FINDINGS")
    print("=" * 60)
    print("""
Key observations:

1. STYLE CAPTURE: The 'Sub-header (CVHP)' style wasn't captured because
   it might not have been used in any paragraph in the document body,
   or the codec didn't encounter it during iteration.

2. RUN MERGING: The codec merges adjacent runs with identical formatting
   which reduces the run count but preserves the text and formatting.
   This is intentional to reduce noise in the JSON.

3. PAGE BREAKS: The codec converts page-break-only paragraphs to
   dedicated page_break blocks. This may cause paragraph index shifts.

4. HEADERS/FOOTERS: These are preserved because the template's section
   properties are maintained (we only replace body content).

5. STYLE DEFINITIONS: The style_defs capture the key properties
   (font, spacing, alignment) which can be used by an LLM to understand
   what each style code means.
""")


if __name__ == "__main__":
    main()
