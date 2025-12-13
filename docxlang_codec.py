"""
docxlang_codec.py

A minimal, LLM-friendly intermediate representation for Word (.docx):

- Converts DOCX -> DocxLang v1 (JSON-serializable dict)
- Converts DocxLang v1 -> DOCX (rendered into a golden template)

Supported blocks:
- paragraphs (with runs + basic run overrides)
- tables (cell paragraphs)
- page breaks

Supported formatting capture:
- paragraph style (S-codes), run style (C-codes), table style (T-codes)
- paragraph overrides: alignment, spacing before/after, indents, keep rules, line spacing
- run overrides: bold/italic/underline/font size/font name/font color

Limitations (by design for v1):
- no images/shapes in-order
- hyperlinks/fields not fully modeled
- arbitrary numbering definitions not round-tripped; prefer list paragraph styles in template
"""

from __future__ import annotations

from typing import Any, Dict, List, Optional, Union

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn, nsmap
from docx.oxml.table import CT_Tbl
from docx.shared import Pt, RGBColor
from docx.table import Table
from docx.text.paragraph import Paragraph


# -----------------------------
# Helpers: formatting conversion
# -----------------------------

def _is_default_char_style(name: Optional[str]) -> bool:
    # python-docx uses "Default Paragraph Font" as the default character style.
    return name is None or name.strip() == "" or name == "Default Paragraph Font"


def _length_to_pt(length) -> Optional[float]:
    if length is None:
        return None
    try:
        return float(length.pt)
    except Exception:
        return None


def _align_to_str(align) -> Optional[str]:
    if align is None:
        return None
    if align == WD_ALIGN_PARAGRAPH.LEFT:
        return "left"
    if align == WD_ALIGN_PARAGRAPH.CENTER:
        return "center"
    if align == WD_ALIGN_PARAGRAPH.RIGHT:
        return "right"
    if align == WD_ALIGN_PARAGRAPH.JUSTIFY:
        return "justify"
    # fallback numeric string for uncommon alignments
    try:
        return str(int(align))
    except Exception:
        return None


def _str_to_align(s: Optional[str]):
    if s is None:
        return None
    if s == "left":
        return WD_ALIGN_PARAGRAPH.LEFT
    if s == "center":
        return WD_ALIGN_PARAGRAPH.CENTER
    if s == "right":
        return WD_ALIGN_PARAGRAPH.RIGHT
    if s == "justify":
        return WD_ALIGN_PARAGRAPH.JUSTIFY
    try:
        return WD_ALIGN_PARAGRAPH(int(s))
    except Exception:
        return None


def _twips_to_pt(twips) -> Optional[float]:
    """Convert twips (1/20 of a point) to points."""
    if twips is None:
        return None
    try:
        return int(twips) / 20.0
    except Exception:
        return None


# -----------------------------
# Numbering/Bullet helpers
# -----------------------------

class NumberingResolver:
    """
    Parses numbering.xml to:
      - map numId -> abstractNumId
      - summarize abstractNumId level defs (numFmt, lvlText, bullet font, indent)
      - ensure a numId exists for a given abstractNumId in the output doc
    """
    def __init__(self, doc: Document):
        self.doc = doc
        self.numbering_el = None
        self.numId_to_abstract: Dict[int, int] = {}
        self.abstract_to_numIds: Dict[int, List[int]] = {}
        self.abstract_levels: Dict[int, Dict[int, Dict[str, Any]]] = {}

        try:
            self.numbering_el = doc.part.numbering_part.element
        except Exception:
            self.numbering_el = None

        if self.numbering_el is not None:
            self._parse_numbering()

    def _parse_numbering(self) -> None:
        numbering = self.numbering_el

        # numId -> abstractNumId
        for num in numbering.iterchildren():
            if num.tag != qn("w:num"):
                continue
            numId_raw = num.get(qn("w:numId"))
            if numId_raw is None:
                continue
            try:
                numId = int(numId_raw)
            except Exception:
                continue

            abs_el = num.find(qn("w:abstractNumId"))
            if abs_el is None:
                continue
            abs_raw = abs_el.get(qn("w:val"))
            if abs_raw is None:
                continue
            try:
                abstractNumId = int(abs_raw)
            except Exception:
                continue

            self.numId_to_abstract[numId] = abstractNumId
            self.abstract_to_numIds.setdefault(abstractNumId, []).append(numId)

        # abstractNumId -> {ilvl -> level summary}
        for absn in numbering.iterchildren():
            if absn.tag != qn("w:abstractNum"):
                continue
            abs_raw = absn.get(qn("w:abstractNumId"))
            if abs_raw is None:
                continue
            try:
                abstractNumId = int(abs_raw)
            except Exception:
                continue

            levels: Dict[int, Dict[str, Any]] = {}
            for lvl in absn.iterchildren():
                if lvl.tag != qn("w:lvl"):
                    continue
                ilvl_raw = lvl.get(qn("w:ilvl"))
                if ilvl_raw is None:
                    continue
                try:
                    ilvl = int(ilvl_raw)
                except Exception:
                    continue

                numFmt_el = lvl.find(qn("w:numFmt"))
                numFmt = numFmt_el.get(qn("w:val")) if numFmt_el is not None else None

                lvlText_el = lvl.find(qn("w:lvlText"))
                lvlText = lvlText_el.get(qn("w:val")) if lvlText_el is not None else None

                # bullet font often stored on lvl/rPr/rFonts
                font = None
                rPr = lvl.find(qn("w:rPr"))
                if rPr is not None:
                    rFonts = rPr.find(qn("w:rFonts"))
                    if rFonts is not None:
                        font = (
                            rFonts.get(qn("w:ascii"))
                            or rFonts.get(qn("w:hAnsi"))
                            or rFonts.get(qn("w:cs"))
                            or rFonts.get(qn("w:hint"))
                        )

                # indent stored on lvl/pPr/ind (twips)
                left_tw = hanging_tw = None
                pPr = lvl.find(qn("w:pPr"))
                if pPr is not None:
                    ind = pPr.find(qn("w:ind"))
                    if ind is not None:
                        left_tw = ind.get(qn("w:left"))
                        hanging_tw = ind.get(qn("w:hanging"))

                levels[ilvl] = {
                    "numFmt": numFmt,
                    "lvlText": lvlText,
                    "font": font,
                    "leftIndentPt": _twips_to_pt(left_tw),
                    "hangingPt": _twips_to_pt(hanging_tw),
                }

            self.abstract_levels[abstractNumId] = levels

    def paragraph_num(self, p: Paragraph) -> Optional[tuple]:
        """Get (numId, ilvl) for a paragraph, or None if not a list item."""
        numId = p._p.xpath("./w:pPr/w:numPr/w:numId/@w:val")
        ilvl = p._p.xpath("./w:pPr/w:numPr/w:ilvl/@w:val")
        if not numId:
            return None
        try:
            n = int(numId[0])
        except Exception:
            return None
        try:
            lvl = int(ilvl[0]) if ilvl else 0
        except Exception:
            lvl = 0
        return (n, lvl)

    def abstract_for_num(self, numId: int) -> Optional[int]:
        return self.numId_to_abstract.get(numId)

    def level_def(self, abstractNumId: int, ilvl: int) -> Optional[Dict[str, Any]]:
        levels = self.abstract_levels.get(abstractNumId) or {}
        return levels.get(ilvl) or levels.get(0)

    def is_bullet(self, numId: int, ilvl: int) -> bool:
        absId = self.abstract_for_num(numId)
        if absId is None:
            return False
        ld = self.level_def(absId, ilvl) or {}
        return (ld.get("numFmt") == "bullet")

    def ensure_numId_for_abstract(self, abstractNumId: int) -> Optional[int]:
        """
        Return an existing numId referencing abstractNumId, or create one in numbering.xml.
        This is only needed at render-time into the template doc.
        """
        if self.numbering_el is None:
            return None

        existing = (self.abstract_to_numIds.get(abstractNumId) or [])
        if existing:
            return existing[0]

        # allocate new numId = max+1
        all_numIds = list(self.numId_to_abstract.keys())
        new_numId = (max(all_numIds) + 1) if all_numIds else 1

        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(new_numId))

        abs_el = OxmlElement("w:abstractNumId")
        abs_el.set(qn("w:val"), str(abstractNumId))
        num.append(abs_el)

        self.numbering_el.append(num)

        self.numId_to_abstract[new_numId] = abstractNumId
        self.abstract_to_numIds.setdefault(abstractNumId, []).append(new_numId)
        return new_numId


def _set_paragraph_numPr(p: Paragraph, numId: int, ilvl: int) -> None:
    """
    Low-level: set w:numPr on paragraph to drive bullets/numbers & nesting level.
    """
    pPr = p._p.get_or_add_pPr()

    # remove existing numPr if present
    existing = pPr.find(qn("w:numPr"))
    if existing is not None:
        pPr.remove(existing)

    numPr = OxmlElement("w:numPr")

    ilvl_el = OxmlElement("w:ilvl")
    ilvl_el.set(qn("w:val"), str(int(ilvl)))
    numPr.append(ilvl_el)

    numId_el = OxmlElement("w:numId")
    numId_el.set(qn("w:val"), str(int(numId)))
    numPr.append(numId_el)

    pPr.append(numPr)


class BulletStyleCatalog:
    """
    Maps abstractNumId -> B-code.
    Stores defs used in DocxLang:
      B1: { abstractNumId: 5, levels: {0:{lvlText:'•',...}, 1:{...}} }
    """
    def __init__(self, base: Optional[Dict[str, Dict[str, Any]]] = None):
        self.code_to_def: Dict[str, Dict[str, Any]] = dict(base or {})
        self.abstract_to_code: Dict[int, str] = {}
        for code, d in self.code_to_def.items():
            absId = d.get("abstractNumId")
            if isinstance(absId, int):
                self.abstract_to_code[absId] = code

    def _next_code(self) -> str:
        max_n = 0
        for code in self.code_to_def.keys():
            if code.startswith("B"):
                try:
                    max_n = max(max_n, int(code[1:]))
                except Exception:
                    pass
        return f"B{max_n + 1}"

    def ensure_from_abstract(self, abstractNumId: int, resolver: NumberingResolver) -> str:
        existing = self.abstract_to_code.get(abstractNumId)
        if existing:
            return existing

        code = self._next_code()
        self.abstract_to_code[abstractNumId] = code

        self.code_to_def[code] = {
            "abstractNumId": abstractNumId,
            "levels": resolver.abstract_levels.get(abstractNumId, {}),
        }
        return code

    def build_from_doc(self, doc: Document, resolver: NumberingResolver, include_unused: bool = True) -> None:
        """
        Include all bullet abstractNum definitions (include_unused=True),
        OR only those actually referenced by paragraphs (include_unused=False).
        """
        bullet_absIds: set = set()

        if include_unused:
            for absId, levels in (resolver.abstract_levels or {}).items():
                # consider it a bullet list if level 0 is bullet OR any level is bullet
                if any((lvl.get("numFmt") == "bullet") for lvl in levels.values()):
                    bullet_absIds.add(absId)
        else:
            for p in doc.paragraphs:
                info = resolver.paragraph_num(p)
                if not info:
                    continue
                numId, ilvl = info
                absId = resolver.abstract_for_num(numId)
                if absId is None:
                    continue
                if resolver.is_bullet(numId, ilvl):
                    bullet_absIds.add(absId)

        for absId in sorted(bullet_absIds):
            self.ensure_from_abstract(absId, resolver)

    def as_dict(self) -> Dict[str, Dict[str, Any]]:
        return dict(self.code_to_def)

    def default_code(self) -> Optional[str]:
        codes = sorted(self.code_to_def.keys(), key=lambda c: int(c[1:]) if c[1:].isdigit() else 999999)
        return codes[0] if codes else None

    def has_multiple(self) -> bool:
        return len(self.code_to_def) > 1


# -----------------------------
# Style catalog (S1/C1/T1 codes)
# -----------------------------

class StyleCatalog:
    """
    Maintains stable codes -> Word style names for:
      - paragraph styles: S1, S2, ...
      - character styles: C1, C2, ...
      - table styles: T1, T2, ...

    If you pass a base catalog built from the golden template, your codes remain stable.
    """
    def __init__(self, styles: Optional[Dict[str, Dict[str, str]]] = None):
        styles = styles or {"paragraph": {}, "character": {}, "table": {}}
        self.code_to_name: Dict[str, Dict[str, str]] = {
            "paragraph": dict(styles.get("paragraph", {})),
            "character": dict(styles.get("character", {})),
            "table": dict(styles.get("table", {})),
        }
        self.name_to_code: Dict[str, Dict[str, str]] = {
            kind: {name: code for code, name in mapping.items()}
            for kind, mapping in self.code_to_name.items()
        }

    def _next_code(self, kind: str) -> str:
        prefix = {"paragraph": "S", "character": "C", "table": "T"}[kind]
        max_n = 0
        for code in self.code_to_name[kind].keys():
            if code.startswith(prefix):
                try:
                    max_n = max(max_n, int(code[len(prefix):]))
                except Exception:
                    pass
        return f"{prefix}{max_n + 1}"

    def ensure(self, kind: str, style_name: Optional[str]) -> Optional[str]:
        if style_name is None:
            return None
        style_name = style_name.strip()
        if style_name == "":
            return None
        if kind == "character" and _is_default_char_style(style_name):
            return None

        existing = self.name_to_code[kind].get(style_name)
        if existing:
            return existing

        code = self._next_code(kind)
        self.code_to_name[kind][code] = style_name
        self.name_to_code[kind][style_name] = code
        return code

    def name(self, kind: str, code: Optional[str]) -> Optional[str]:
        if not code:
            return None
        return self.code_to_name.get(kind, {}).get(code)

    def as_dict(self) -> Dict[str, Dict[str, str]]:
        return {k: dict(v) for k, v in self.code_to_name.items()}


# -----------------------------
# Iterate DOCX body in order
# -----------------------------

def iter_block_items(doc: Document):
    """
    Yield Paragraph and Table objects in document order (body only).
    """
    body = doc.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, doc)
        elif child.tag == qn("w:tbl"):
            yield Table(child, doc)


# -----------------------------
# DOCX -> DocxLang
# -----------------------------

def _paragraph_contains_page_break(p: Paragraph) -> bool:
    return len(p._p.xpath(".//w:br[@w:type='page']")) > 0


def _get_numbering_info(p: Paragraph) -> Optional[Dict[str, Any]]:
    """
    Minimal numbering info (numId, level) if present.
    NOTE: This does not resolve bullet vs number from numbering.xml; we store a heuristic "kind".
    """
    numId = p._p.xpath("./w:pPr/w:numPr/w:numId/@w:val")
    ilvl = p._p.xpath("./w:pPr/w:numPr/w:ilvl/@w:val")
    if not numId and not ilvl:
        return None

    info: Dict[str, Any] = {}
    if numId:
        try:
            info["numId"] = int(numId[0])
        except Exception:
            info["numId"] = numId[0]
    if ilvl:
        try:
            info["level"] = int(ilvl[0])
        except Exception:
            info["level"] = ilvl[0]

    style_name = (p.style.name if p.style is not None else "") or ""
    if "bullet" in style_name.lower():
        info["kind"] = "bullet"
    elif "number" in style_name.lower():
        info["kind"] = "number"
    else:
        info["kind"] = "unknown"
    return info


def paragraph_to_block(
    p: Paragraph,
    catalog: StyleCatalog,
    resolver: Optional[NumberingResolver] = None,
    bullet_styles: Optional[BulletStyleCatalog] = None,
) -> Dict[str, Any]:
    pStyle = catalog.ensure("paragraph", p.style.name if p.style is not None else None)

    override: Dict[str, Any] = {}
    a = _align_to_str(p.alignment)
    if a:
        override["align"] = a

    pf = p.paragraph_format
    for key, val in [
        ("spaceBeforePt", _length_to_pt(pf.space_before)),
        ("spaceAfterPt", _length_to_pt(pf.space_after)),
        ("leftIndentPt", _length_to_pt(pf.left_indent)),
        ("rightIndentPt", _length_to_pt(pf.right_indent)),
        ("firstLineIndentPt", _length_to_pt(pf.first_line_indent)),
    ]:
        if val is not None:
            override[key] = val

    if pf.keep_with_next is not None:
        override["keepWithNext"] = bool(pf.keep_with_next)
    if pf.keep_together is not None:
        override["keepLinesTogether"] = bool(pf.keep_together)
    if pf.page_break_before is not None:
        override["pageBreakBefore"] = bool(pf.page_break_before)

    if pf.line_spacing is not None:
        if isinstance(pf.line_spacing, (int, float)):
            override["lineSpacing"] = float(pf.line_spacing)
            override["lineSpacingRule"] = "multiple"
        else:
            try:
                override["lineSpacing"] = float(pf.line_spacing.pt)
                override["lineSpacingRule"] = "exact"
            except Exception:
                pass

    # List/bullet capture driven by numbering.xml
    if resolver is not None and bullet_styles is not None:
        numinfo = resolver.paragraph_num(p)
        if numinfo:
            numId, ilvl = numinfo
            absId = resolver.abstract_for_num(numId)
            if absId is not None and resolver.is_bullet(numId, ilvl):
                bcode = bullet_styles.ensure_from_abstract(absId, resolver)

                # store LLM-friendly list ref (do NOT store raw numId)
                override["list"] = {"kind": "bullet", "level": ilvl}
                # only include explicit style if there are multiple
                if bullet_styles.has_multiple():
                    override["list"]["style"] = bcode

                # avoid fighting list indentation: let numbering definition control it
                override.pop("leftIndentPt", None)
                override.pop("firstLineIndentPt", None)
                override.pop("rightIndentPt", None)
    else:
        # Fallback to old numbering info if resolver not provided
        numinfo = _get_numbering_info(p)
        if numinfo:
            override["numbering"] = numinfo

    runs: List[Dict[str, Any]] = []
    for r in p.runs:
        has_page_break = len(r._r.xpath(".//w:br[@w:type='page']")) > 0
        if has_page_break and (r.text or "") == "":
            runs.append({"type": "break", "breakType": "page"})
            continue

        text = r.text or ""
        if text == "" and not has_page_break:
            continue

        rStyleName = None
        try:
            rStyleName = r.style.name if r.style is not None else None
        except Exception:
            rStyleName = None
        rStyle = catalog.ensure("character", rStyleName)

        r_override: Dict[str, Any] = {}
        if r.bold is not None:
            r_override["bold"] = bool(r.bold)
        if r.italic is not None:
            r_override["italic"] = bool(r.italic)
        if r.underline is not None:
            r_override["underline"] = bool(r.underline)

        f = r.font
        if f is not None:
            fs = _length_to_pt(f.size)
            if fs is not None:
                r_override["fontSizePt"] = fs
            if f.name:
                r_override["fontName"] = f.name
            try:
                if f.color and f.color.rgb:
                    r_override["color"] = str(f.color.rgb)  # e.g. 'FF0000'
            except Exception:
                pass

        run_item: Dict[str, Any] = {"type": "text", "text": text}
        if rStyle:
            run_item["rStyle"] = rStyle
        if r_override:
            run_item["override"] = r_override
        runs.append(run_item)

    # Merge adjacent text runs with identical formatting to reduce noise
    merged: List[Dict[str, Any]] = []
    for item in runs:
        if (
            merged
            and item["type"] == "text"
            and merged[-1]["type"] == "text"
            and item.get("rStyle") == merged[-1].get("rStyle")
            and item.get("override") == merged[-1].get("override")
        ):
            merged[-1]["text"] += item.get("text", "")
        else:
            merged.append(item)

    out: Dict[str, Any] = {"type": "paragraph", "runs": merged}
    if pStyle:
        out["pStyle"] = pStyle
    if override:
        out["override"] = override
    return out


def table_to_block(
    t: Table,
    catalog: StyleCatalog,
    resolver: Optional[NumberingResolver] = None,
    bullet_styles: Optional[BulletStyleCatalog] = None,
) -> Dict[str, Any]:
    tblStyleName = None
    try:
        tblStyleName = t.style.name if t.style is not None else None
    except Exception:
        pass
    tblStyle = catalog.ensure("table", tblStyleName)

    rows_data: List[List[Dict[str, Any]]] = []
    for row in t.rows:
        row_cells: List[Dict[str, Any]] = []
        for cell in row.cells:
            cell_blocks: List[Dict[str, Any]] = []
            for p in cell.paragraphs:
                if (p.text or "").strip() == "" and len(p.runs) == 0:
                    continue
                cell_blocks.append(paragraph_to_block(p, catalog, resolver, bullet_styles))
            if not cell_blocks:
                cell_blocks.append({"type": "paragraph", "runs": []})
            row_cells.append({"blocks": cell_blocks})
        rows_data.append(row_cells)

    out: Dict[str, Any] = {"type": "table", "rows": rows_data}
    if tblStyle:
        out["tblStyle"] = tblStyle
    return out


def build_style_catalog_from_docx(
    docx_path: str,
    include_unused: bool = False,
    include_builtin: bool = True,
) -> Dict[str, Dict[str, str]]:
    """
    Build a stable style catalog (codes -> style names) from a DOCX file.

    - include_unused=True: include all styles in the doc (filtered by builtin/type)
      (useful for a "template style pack")
    - include_unused=False: include only styles actually referenced in body content
    """
    doc = Document(docx_path)

    used_par, used_char, used_tbl = set(), set(), set()

    if include_unused:
        for st in doc.styles:
            if not include_builtin and getattr(st, "builtin", False):
                continue
            if st.type == WD_STYLE_TYPE.PARAGRAPH:
                used_par.add(st.name)
            elif st.type == WD_STYLE_TYPE.CHARACTER:
                if not _is_default_char_style(st.name):
                    used_char.add(st.name)
            elif st.type == WD_STYLE_TYPE.TABLE:
                used_tbl.add(st.name)
    else:
        for item in iter_block_items(doc):
            if isinstance(item, Paragraph):
                if item.style is not None:
                    used_par.add(item.style.name)
                for r in item.runs:
                    try:
                        if r.style is not None and not _is_default_char_style(r.style.name):
                            used_char.add(r.style.name)
                    except Exception:
                        pass
            elif isinstance(item, Table):
                try:
                    if item.style is not None:
                        used_tbl.add(item.style.name)
                except Exception:
                    pass
                for row in item.rows:
                    for cell in row.cells:
                        for p in cell.paragraphs:
                            if p.style is not None:
                                used_par.add(p.style.name)
                            for r in p.runs:
                                try:
                                    if r.style is not None and not _is_default_char_style(r.style.name):
                                        used_char.add(r.style.name)
                                except Exception:
                                    pass

    def assign(names, prefix):
        return {f"{prefix}{i}": name for i, name in enumerate(sorted(names), start=1)}

    return {
        "paragraph": assign(used_par, "S"),
        "character": assign(used_char, "C"),
        "table": assign(used_tbl, "T"),
    }


# Optional: style summaries for LLM friendliness

def _summarize_font(font) -> Dict[str, Any]:
    out: Dict[str, Any] = {}
    if getattr(font, "name", None):
        out["name"] = font.name
    if getattr(font, "size", None) is not None:
        try:
            out["sizePt"] = float(font.size.pt)
        except Exception:
            pass
    if getattr(font, "bold", None) is not None:
        out["bold"] = bool(font.bold)
    if getattr(font, "italic", None) is not None:
        out["italic"] = bool(font.italic)
    if getattr(font, "underline", None) is not None:
        out["underline"] = bool(font.underline)
    try:
        if font.color and font.color.rgb:
            out["color"] = str(font.color.rgb)
    except Exception:
        pass
    return out


def _summarize_paragraph_format(pf) -> Dict[str, Any]:
    out: Dict[str, Any] = {}
    a = _align_to_str(getattr(pf, "alignment", None))
    if a:
        out["align"] = a
    for key, val in [
        ("spaceBeforePt", _length_to_pt(getattr(pf, "space_before", None))),
        ("spaceAfterPt", _length_to_pt(getattr(pf, "space_after", None))),
        ("leftIndentPt", _length_to_pt(getattr(pf, "left_indent", None))),
        ("rightIndentPt", _length_to_pt(getattr(pf, "right_indent", None))),
        ("firstLineIndentPt", _length_to_pt(getattr(pf, "first_line_indent", None))),
    ]:
        if val is not None:
            out[key] = val
    if getattr(pf, "keep_with_next", None) is not None:
        out["keepWithNext"] = bool(pf.keep_with_next)
    if getattr(pf, "keep_together", None) is not None:
        out["keepLinesTogether"] = bool(pf.keep_together)
    if getattr(pf, "page_break_before", None) is not None:
        out["pageBreakBefore"] = bool(pf.page_break_before)

    ls = getattr(pf, "line_spacing", None)
    if ls is not None:
        if isinstance(ls, (int, float)):
            out["lineSpacing"] = float(ls)
            out["lineSpacingRule"] = "multiple"
        else:
            try:
                out["lineSpacing"] = float(ls.pt)
                out["lineSpacingRule"] = "exact"
            except Exception:
                pass
    return out


def _summarize_style(doc: Document, style_name: str) -> Dict[str, Any]:
    out: Dict[str, Any] = {"name": style_name}
    try:
        st = doc.styles[style_name]
    except KeyError:
        return out

    try:
        if st.base_style is not None:
            out["base"] = st.base_style.name
    except Exception:
        pass

    try:
        fsum = _summarize_font(st.font)
        if fsum:
            out["font"] = fsum
    except Exception:
        pass

    try:
        psum = _summarize_paragraph_format(st.paragraph_format)
        if psum:
            out["para"] = psum
    except Exception:
        pass

    try:
        out["type"] = str(st.type)  # e.g. 'PARAGRAPH (1)'
    except Exception:
        pass

    return out


def extract_style_defs(doc: Document, catalog: StyleCatalog) -> Dict[str, Dict[str, Dict[str, Any]]]:
    defs: Dict[str, Dict[str, Dict[str, Any]]] = {"paragraph": {}, "character": {}, "table": {}}
    for kind in ["paragraph", "character", "table"]:
        for code, name in catalog.code_to_name[kind].items():
            defs[kind][code] = _summarize_style(doc, name)
    return defs


def docx_to_docxlang(
    docx_path: str,
    base_catalog: Optional[Dict[str, Dict[str, str]]] = None,
    base_list_styles: Optional[Dict[str, Any]] = None,
    include_style_defs: bool = False,
    group_bullets: bool = True,
) -> Dict[str, Any]:
    """
    Convert a .docx to DocxLang v1.

    If base_catalog is provided (recommended: derived from your golden template),
    codes remain stable and the catalog is reused/extended.

    If group_bullets is True (default), consecutive bullet paragraphs are grouped
    into bulleted_list blocks with nested level information.
    """
    doc = Document(docx_path)
    catalog = StyleCatalog(base_catalog)

    # Set up numbering/bullet handling
    resolver = NumberingResolver(doc)
    bullet_styles = BulletStyleCatalog(
        (base_list_styles or {}).get("bullets") if base_list_styles else None
    )
    bullet_styles.build_from_doc(doc, resolver, include_unused=True)

    blocks: List[Dict[str, Any]] = []
    pending_bullets: Optional[Dict[str, Any]] = None

    def flush_pending_bullets():
        nonlocal pending_bullets
        if pending_bullets is not None:
            blocks.append(pending_bullets)
            pending_bullets = None

    for item in iter_block_items(doc):
        if isinstance(item, Paragraph):
            # If it's a pure page-break paragraph, represent it as a block.
            if _paragraph_contains_page_break(item) and (item.text or "").strip() == "":
                flush_pending_bullets()
                pb_block: Dict[str, Any] = {"type": "page_break"}
                # Preserve paragraph style on page break blocks
                if item.style is not None:
                    pStyle = catalog.ensure("paragraph", item.style.name)
                    if pStyle:
                        pb_block["pStyle"] = pStyle
                blocks.append(pb_block)
            else:
                pb = paragraph_to_block(item, catalog, resolver, bullet_styles)
                list_ref = (pb.get("override") or {}).get("list")

                if group_bullets and list_ref and list_ref.get("kind") == "bullet":
                    # Extract list info from paragraph for grouping
                    level = int(list_ref.get("level", 0))
                    style = list_ref.get("style")

                    # Remove list from paragraph override (list block owns it)
                    if pb.get("override"):
                        pb["override"].pop("list", None)
                        if not pb["override"]:
                            pb.pop("override", None)

                    # Start new list or continue existing one
                    if pending_bullets is None or pending_bullets.get("style") != style:
                        flush_pending_bullets()
                        pending_bullets = {"type": "bulleted_list", "items": []}
                        if style is not None:
                            pending_bullets["style"] = style

                    pending_bullets["items"].append({"level": level, "paragraph": pb})
                else:
                    # Non-bullet paragraph
                    flush_pending_bullets()
                    blocks.append(pb)
        elif isinstance(item, Table):
            flush_pending_bullets()
            blocks.append(table_to_block(item, catalog, resolver, bullet_styles))

    # Flush any remaining bullets
    flush_pending_bullets()

    out: Dict[str, Any] = {
        "schema": "docxlang/v1",
        "styles": catalog.as_dict(),
        "blocks": blocks,
    }

    # Add list styles if any bullet definitions exist
    if bullet_styles.as_dict():
        out["list_styles"] = {"bullets": bullet_styles.as_dict()}
        default_b = bullet_styles.default_code()
        if default_b:
            out["defaults"] = {"bullet": default_b}

    if include_style_defs:
        out["style_defs"] = extract_style_defs(doc, catalog)
    return out


# -----------------------------
# DocxLang -> DOCX (template render)
# -----------------------------

def _clear_paragraph(p: Paragraph) -> None:
    # Remove all run XML nodes
    for r in list(p.runs):
        try:
            r._r.getparent().remove(r._r)
        except Exception:
            pass


def _apply_paragraph_override(p: Paragraph, override: Dict[str, Any]) -> None:
    if not override:
        return
    if "align" in override:
        p.alignment = _str_to_align(override["align"])

    pf = p.paragraph_format
    for key, attr in [
        ("spaceBeforePt", "space_before"),
        ("spaceAfterPt", "space_after"),
        ("leftIndentPt", "left_indent"),
        ("rightIndentPt", "right_indent"),
        ("firstLineIndentPt", "first_line_indent"),
    ]:
        if key in override:
            setattr(pf, attr, Pt(float(override[key])))

    if "keepWithNext" in override:
        pf.keep_with_next = bool(override["keepWithNext"])
    if "keepLinesTogether" in override:
        pf.keep_together = bool(override["keepLinesTogether"])
    if "pageBreakBefore" in override:
        pf.page_break_before = bool(override["pageBreakBefore"])
    if "lineSpacing" in override:
        pf.line_spacing = override["lineSpacing"]


def _apply_run_override(run, override: Dict[str, Any]) -> None:
    if not override:
        return
    if "bold" in override:
        run.bold = bool(override["bold"])
    if "italic" in override:
        run.italic = bool(override["italic"])
    if "underline" in override:
        run.underline = bool(override["underline"])

    f = run.font
    if "fontSizePt" in override:
        f.size = Pt(float(override["fontSizePt"]))
    if "fontName" in override:
        f.name = str(override["fontName"])
    if "color" in override:
        try:
            rgb = str(override["color"]).replace("#", "")
            f.color.rgb = RGBColor.from_string(rgb)
        except Exception:
            pass


def _apply_list_ref(
    p: Paragraph,
    list_ref: Dict[str, Any],
    bullet_defs: Dict[str, Dict[str, Any]],
    defaults: Dict[str, Any],
    num_resolver: NumberingResolver,
) -> None:
    """Apply bullet/numbering to a paragraph based on list_ref."""
    if not list_ref:
        return
    if list_ref.get("kind") != "bullet":
        return

    level = int(list_ref.get("level", 0))
    style = list_ref.get("style") or defaults.get("bullet")
    if not style:
        return
    bdef = bullet_defs.get(style)
    if not bdef:
        return

    absId = bdef.get("abstractNumId")
    if not isinstance(absId, int):
        return

    numId = num_resolver.ensure_numId_for_abstract(absId)
    if numId is None:
        return

    _set_paragraph_numPr(p, numId=numId, ilvl=level)


def _render_paragraph(
    p: Paragraph,
    block: Dict[str, Any],
    catalog: StyleCatalog,
    bullet_defs: Optional[Dict[str, Dict[str, Any]]] = None,
    defaults: Optional[Dict[str, Any]] = None,
    num_resolver: Optional[NumberingResolver] = None,
) -> None:
    _clear_paragraph(p)

    style_name = catalog.name("paragraph", block.get("pStyle"))
    if style_name:
        try:
            p.style = style_name
        except KeyError:
            pass

    _apply_paragraph_override(p, block.get("override") or {})

    for run_item in block.get("runs") or []:
        if run_item.get("type") == "break":
            if run_item.get("breakType") == "page":
                p.add_run().add_break(WD_BREAK.PAGE)
            else:
                p.add_run().add_break()
            continue

        run = p.add_run(run_item.get("text", ""))
        r_style_name = catalog.name("character", run_item.get("rStyle"))
        if r_style_name:
            try:
                run.style = r_style_name
            except KeyError:
                pass
        _apply_run_override(run, run_item.get("override") or {})

    # Apply bullets/nesting after content is in place
    if bullet_defs is not None and defaults is not None and num_resolver is not None:
        list_ref = (block.get("override") or {}).get("list")
        if list_ref:
            _apply_list_ref(p, list_ref, bullet_defs, defaults, num_resolver)


def _insert_table_before(anchor: Paragraph, rows: int, cols: int, doc: Document) -> Table:
    """
    Insert a table XML element directly before `anchor` and return a python-docx Table wrapper.
    """
    section = doc.sections[0]
    try:
        usable_width = section.page_width - section.left_margin - section.right_margin
    except Exception:
        # fallback width if margins aren't available
        from docx.shared import Inches
        usable_width = Inches(6)

    tbl = CT_Tbl.new_tbl(rows, cols, usable_width)
    anchor._p.addprevious(tbl)
    return Table(tbl, anchor._parent)


def _render_table_before(
    anchor: Paragraph,
    block: Dict[str, Any],
    catalog: StyleCatalog,
    doc: Document,
    bullet_defs: Optional[Dict[str, Dict[str, Any]]] = None,
    defaults: Optional[Dict[str, Any]] = None,
    num_resolver: Optional[NumberingResolver] = None,
) -> None:
    rows_data = block.get("rows") or []
    rows = len(rows_data) or 1
    cols = len(rows_data[0]) if rows_data else 1

    tbl = _insert_table_before(anchor, rows, cols, doc)

    tbl_style_name = catalog.name("table", block.get("tblStyle"))
    if tbl_style_name:
        try:
            tbl.style = tbl_style_name
        except KeyError:
            pass

    for r_i in range(rows):
        for c_i in range(cols):
            cell_obj = (
                rows_data[r_i][c_i]
                if rows_data and r_i < len(rows_data) and c_i < len(rows_data[r_i])
                else {"blocks": []}
            )
            cell = tbl.cell(r_i, c_i)
            try:
                # clears all cell content including the default empty paragraph
                cell._tc.clear_content()
            except Exception:
                pass

            for para_block in (cell_obj.get("blocks") or []):
                if para_block.get("type") != "paragraph":
                    continue
                cp = cell.add_paragraph()
                _render_paragraph(cp, para_block, catalog, bullet_defs, defaults, num_resolver)


def _find_anchor_paragraph(doc: Document, anchor_text: str) -> Optional[Paragraph]:
    # Search body paragraphs
    for p in doc.paragraphs:
        if anchor_text in (p.text or ""):
            return p
    # Also search table cell paragraphs (in case placeholder is placed in a table)
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    if anchor_text in (p.text or ""):
                        return p
    return None


def docxlang_to_docx(
    docxlang: Dict[str, Any],
    template_path: str,
    output_path: str,
    anchor_text: str = "<<CONTENT>>",
    clear_body_if_no_anchor: bool = False,
) -> None:
    """
    Render DocxLang v1 into a .docx using a golden template.

    Strategy:
    - If anchor_text exists: insert blocks *before* anchor, then remove anchor.
      (This preserves anything before/after the placeholder and preserves headers/footers.)
    - If anchor_text not found:
        - optionally clear body content (preserving section properties) if clear_body_if_no_anchor=True
        - otherwise append blocks at the end
    """
    doc = Document(template_path)
    catalog = StyleCatalog(docxlang.get("styles") or {})
    blocks = docxlang.get("blocks") or []

    # Set up bullet/list handling
    defaults = docxlang.get("defaults") or {}
    bullet_defs = (docxlang.get("list_styles") or {}).get("bullets") or {}
    num_resolver = NumberingResolver(doc)

    anchor = _find_anchor_paragraph(doc, anchor_text)

    def render_block_append(b: Dict[str, Any]) -> None:
        """Render a block by appending to end of document."""
        btype = b.get("type")
        if btype == "page_break":
            p = doc.add_paragraph()
            p.add_run().add_break(WD_BREAK.PAGE)
            style_name = catalog.name("paragraph", b.get("pStyle"))
            if style_name:
                try:
                    p.style = style_name
                except KeyError:
                    pass
        elif btype == "paragraph":
            p = doc.add_paragraph()
            _render_paragraph(p, b, catalog, bullet_defs, defaults, num_resolver)
        elif btype == "bulleted_list":
            # Render each item in the bulleted list
            style = b.get("style")
            for it in (b.get("items") or []):
                level = int(it.get("level", 0))
                para_block = it.get("paragraph") or {"type": "paragraph", "runs": []}
                p = doc.add_paragraph()
                _render_paragraph(p, para_block, catalog, bullet_defs, defaults, num_resolver)
                # Apply the correct bullet scheme + level
                list_ref = {"kind": "bullet", "level": level}
                if style:
                    list_ref["style"] = style
                _apply_list_ref(p, list_ref, bullet_defs, defaults, num_resolver)
        elif btype == "table":
            rows_data = b.get("rows") or []
            rows = len(rows_data) or 1
            cols = len(rows_data[0]) if rows_data else 1
            tbl = doc.add_table(rows=rows, cols=cols)
            tbl_style_name = catalog.name("table", b.get("tblStyle"))
            if tbl_style_name:
                try:
                    tbl.style = tbl_style_name
                except KeyError:
                    pass
            for r_i in range(rows):
                for c_i in range(cols):
                    cell_obj = (
                        rows_data[r_i][c_i]
                        if rows_data and r_i < len(rows_data) and c_i < len(rows_data[r_i])
                        else {"blocks": []}
                    )
                    cell = tbl.cell(r_i, c_i)
                    try:
                        cell._tc.clear_content()
                    except Exception:
                        pass
                    for para_block in (cell_obj.get("blocks") or []):
                        if para_block.get("type") != "paragraph":
                            continue
                        cp = cell.add_paragraph()
                        _render_paragraph(cp, para_block, catalog, bullet_defs, defaults, num_resolver)

    def render_block_before_anchor(b: Dict[str, Any]) -> None:
        """Render a block by inserting before anchor paragraph."""
        btype = b.get("type")
        if btype == "page_break":
            p = anchor.insert_paragraph_before()
            p.add_run().add_break(WD_BREAK.PAGE)
            style_name = catalog.name("paragraph", b.get("pStyle"))
            if style_name:
                try:
                    p.style = style_name
                except KeyError:
                    pass
        elif btype == "paragraph":
            p = anchor.insert_paragraph_before()
            _render_paragraph(p, b, catalog, bullet_defs, defaults, num_resolver)
        elif btype == "bulleted_list":
            # Render each item in the bulleted list
            style = b.get("style")
            for it in (b.get("items") or []):
                level = int(it.get("level", 0))
                para_block = it.get("paragraph") or {"type": "paragraph", "runs": []}
                p = anchor.insert_paragraph_before()
                _render_paragraph(p, para_block, catalog, bullet_defs, defaults, num_resolver)
                # Apply the correct bullet scheme + level
                list_ref = {"kind": "bullet", "level": level}
                if style:
                    list_ref["style"] = style
                _apply_list_ref(p, list_ref, bullet_defs, defaults, num_resolver)
        elif btype == "table":
            _render_table_before(anchor, b, catalog, doc, bullet_defs, defaults, num_resolver)

    if anchor is None:
        if clear_body_if_no_anchor:
            doc._body.clear_content()

        # Append blocks at end in forward order
        for b in blocks:
            render_block_append(b)

        doc.save(output_path)
        return

    # Insert blocks before anchor in forward order (insert_paragraph_before inserts just above)
    for b in blocks:
        render_block_before_anchor(b)

    # Remove the anchor paragraph itself
    try:
        anchor._p.getparent().remove(anchor._p)
    except Exception:
        pass

    doc.save(output_path)


# -----------------------------
# Example usage
# -----------------------------
if __name__ == "__main__":
    import json

    # 1) Build a stable style catalog from the golden template (recommended)
    # style_catalog = build_style_catalog_from_docx(
    #     "golden_template.docx",
    #     include_unused=True,   # template often has few body refs; include full style library
    #     include_builtin=False  # usually keep just customer custom styles (tune as needed)
    # )

    # 2) Convert a sample doc to DocxLang using that base catalog
    # docxlang = docx_to_docxlang("sample.docx", base_catalog=style_catalog, include_style_defs=True)
    # with open("sample.docxlang.json", "w", encoding="utf-8") as f:
    #     json.dump(docxlang, f, indent=2, ensure_ascii=False)

    # 3) LLM produces a modified docxlang.json (same schema)
    # 4) Render into the golden template
    # with open("llm_output.docxlang.json", "r", encoding="utf-8") as f:
    #     llm_doc = json.load(f)
    # docxlang_to_docx(llm_doc, "golden_template.docx", "final_output.docx", anchor_text="<<CONTENT>>")
    pass
